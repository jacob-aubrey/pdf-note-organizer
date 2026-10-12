"""Regressions found during the second audit: readability and failure handling."""
import json
from pathlib import Path
import tempfile
import time
import tkinter as tk
from unittest.mock import patch, Mock
import unittest

from watchdog.events import FileMovedEvent, DirCreatedEvent

from note_organizer.exporting import build_notebook, paper_title
from note_organizer.service import EventHandler, Service
from note_organizer.storage import Store, DEFAULT_SETTINGS, entry_sort_key, matches_search, SEARCH_FIELDS
from note_organizer.extraction import extract
from note_organizer.ui import App
from tests.test_notebook import sample_pdf
from tests.ui_helpers import settle_library


def settle_notes(app, root):
    """Let the actual asynchronous reader complete before asserting its result."""
    deadline = time.monotonic() + 3
    while app.note_count.get() == "Loading notes…" and time.monotonic() < deadline:
        app.poll()
        root.update()
        time.sleep(.005)
    if app.note_count.get() == "Loading notes…":
        raise AssertionError("Reader did not finish loading")


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.pdf = self.folder / "paper.pdf"
        sample_pdf(self.pdf)
        self.store = Store(self.folder / "library.sqlite3")
        self.pid = self.store.add_project(self.folder)
        self.service = Service(self.store)
        self.service.scan(self.store.project(self.pid), self.pdf)

    def tearDown(self):
        self.temp.cleanup()

    def test_reading_layout_has_one_page_heading_and_no_export_chatter(self):
        entries = list(reversed(self.store.entries()))
        doc = build_notebook("Research notes", entries)
        pages = [p for p in doc.paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(len(pages), 1)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("Highlighted text:", text)
        self.assertNotIn("Original page with annotations", text)
        self.assertNotIn("Generated", text)
        self.assertNotIn(str(self.folder), text)
        self.assertIn("remained stable", text)
        self.assertIn("Check this claim", text)
        self.assertEqual(len(doc.inline_shapes), 2)

    def test_detailed_layout_remains_available(self):
        doc = build_notebook("Notes", self.store.entries(), layout="Detailed")
        self.assertIn("Highlighted text:", " ".join(p.text for p in doc.paragraphs))

    def test_two_columns_follow_text_flow_instead_of_vertical_position(self):
        import fitz
        pdf = self.folder / "columns.pdf"
        with fitz.open() as source:
            page = source.new_page()
            page.insert_text((50, 650), "First passage in the left column.")
            page.insert_text((320, 80), "Second passage in the right column.")
            for phrase in ("First passage", "Second passage"):
                page.add_highlight_annot(page.search_for(phrase))
            source.save(pdf)
        _, _, meta, entries = extract(pdf, DEFAULT_SETTINGS)
        for e in entries:
            e.update(path=str(pdf), metadata=meta)
        ordered = sorted(entries, key=entry_sort_key)
        self.assertEqual([e["quote"] for e in ordered], ["First passage", "Second passage"])
        self.assertGreater(ordered[0]["rect"][1], ordered[1]["rect"][1])
        doc = build_notebook("Notes", entries)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "Quote"],
                         ["First passage", "Second passage"])

    def test_layout_change_does_not_repeat_extraction(self):
        self.store.set_settings(self.pid, {**DEFAULT_SETTINGS, "notebook_style": "Detailed"})
        with patch("note_organizer.service.extract") as reader:
            self.service.scan(self.store.project(self.pid), self.pdf)
        reader.assert_not_called()
        self.assertTrue(self.store.project(self.pid)["dirty"])

    def test_extractor_upgrade_refreshes_unchanged_pdf(self):
        with patch("note_organizer.extraction.EXTRACTION_REVISION", 999):
            with patch("note_organizer.service.extract", wraps=extract) as reader:
                self.service.scan(self.store.project(self.pid), self.pdf)
        reader.assert_called_once()

    def test_typesetter_title_falls_back_to_filename_and_quotes_are_not_italic(self):
        self.assertEqual(paper_title("Jost_2023.pdf", {"title": "IR_230225 515..522"}), "Jost 2023")
        self.assertEqual(paper_title("paper.pdf", {"title": "A real title"}), "A real title")
        self.assertFalse(build_notebook("Notes", self.store.entries()).styles["Quote"].font.italic)

    def test_duplicate_page_context_image_appears_once_without_losing_comments(self):
        e = next(e for e in self.store.entries() if e["kind"] == "drawing")
        doc = build_notebook("Notes", [e, {**e, "key": "different", "comment": "Another observation"}])
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertIn("Another observation", " ".join(p.text for p in doc.paragraphs))

    def test_xml_control_character_does_not_break_export(self):
        e = self.store.entries()[0]
        original = e["comment"] = "Read\x00this\x0bnote"
        doc = build_notebook("Notes", [e])
        doc.save(self.folder / "clean.docx")
        self.assertEqual(e["comment"], original)

    def test_word_export_move_event_does_not_trigger_pdf_rescan(self):
        handler = EventHandler(self.service)
        handler.on_moved(FileMovedEvent("temporary.docx", "notebook.docx"))
        self.assertFalse(self.service.sync_requested)
        self.assertFalse(self.service.pending)
        handler.on_moved(FileMovedEvent("temporary.tmp", str(self.pdf)))
        self.assertTrue(self.service.sync_requested)
        self.assertTrue(self.service.pending)

    def test_new_directory_requests_discovery(self):
        EventHandler(self.service).on_created(DirCreatedEvent(str(self.folder / "new")))
        self.assertTrue(self.service.sync_requested)

    def test_repeated_locked_export_preserves_only_one_backup_revision(self):
        self.service.export_pending()
        from docx import Document
        output = self.store.project(self.pid)["output"]
        doc = Document(output)
        doc.add_paragraph("User writing")
        doc.save(output)
        self.store.mark_dirty(self.pid)
        with self.assertLogs(level="ERROR"):
            with patch("note_organizer.exporting.publish_staged", side_effect=PermissionError("Word lock")):
                self.service.export_pending()
                self.service.export_pending()
        backups = list((self.folder / "Notebook backups").glob("*.docx"))
        self.assertEqual(len(backups), 1)
        self.assertIn("User writing", " ".join(p.text for p in Document(backups[0]).paragraphs))

    def test_scan_can_cancel_without_replacing_saved_notes(self):
        with self.assertRaises(InterruptedError):
            extract(self.pdf, DEFAULT_SETTINGS, cancelled=lambda: True)
        self.assertEqual(len(self.store.entries()), 3)

    def test_sync_summary_reports_export_waiting(self):
        with self.assertLogs(level="ERROR"):
            with patch("note_organizer.exporting.publish_staged", side_effect=PermissionError("Word lock")):
                self.service.sync()
        messages = []
        while not self.service.messages.empty():
            messages.append(self.service.messages.get_nowait()[0])
        self.assertIn("notebooks waiting to save", messages[-1])

    def test_search_matches_printed_context_for_handwriting(self):
        self.assertTrue(self.store.entries(query="control experiment"))

    def test_search_fields_do_not_leak_context_into_handwriting(self):
        self.assertTrue(self.store.entries(query="control experiment", fields=["context"]))
        self.assertFalse(self.store.entries(query="control experiment", fields=["ocr"]))
        self.assertTrue(self.store.entries(query="Check this claim", fields=["comment"]))
        self.assertFalse(self.store.entries(query="Check this claim", fields=["quote"]))

    def test_search_words_phrase_and_type_filters(self):
        self.assertTrue(self.store.entries(query="stable remained", fields=["quote"]))
        self.assertFalse(self.store.entries(query="stable remained", fields=["quote"], match_mode="Exact phrase"))
        self.assertTrue(self.store.entries(query="  remained   stable  ", fields=["quote"], match_mode="Exact phrase"))
        self.assertEqual(len(self.store.entries(note_type="Drawings and handwriting")), 2)
        self.assertEqual(len(self.store.entries(note_type="Highlights and markup")), 1)
        self.assertEqual(len(self.store.entries(note_type="Typed notes")), 1)

    def test_bibliographic_fields_and_ocr_are_separate(self):
        entry = {"path": "paper.pdf", "kind": "drawing", "ocr": "Check the concentration",
                 "metadata": {"title": "Contrast imaging", "author": "Jane Scientist", "doi": "10.1234/example"},
                 "page": 2, "label": "iv", "color": "Blue"}
        for field, query in [("title", "contrast"), ("author", "scientist"), ("doi", "10.1234"),
                             ("filename", "paper.pdf"), ("page", "iv"), ("color", "blue"), ("ocr", "concentration")]:
            self.assertTrue(matches_search(entry, query, [field], "All words", "All notes"))
            self.assertFalse(matches_search(entry, query, [k for k in SEARCH_FIELDS if k != field], "All words", "All notes"))

    def test_library_reader_selection_survives_background_refresh(self):
        root = tk.Tk()
        root.withdraw()
        errors = []
        root.report_callback_exception = lambda *args: errors.append(args)
        try:
            with patch.object(Service, "start"):
                app = App(root, self.store, self.folder)
                settle_library(app, root)
                root.update()
                self.assertEqual([app.tabs.tab(t, "text") for t in app.tabs.tabs()], ["Inbox", "Library", "Notes", "reMarkable"])
                paper = app.papers.get_children()[0]
                app.papers.selection_set(paper)
                app.view_notes()
                root.update()
                self.assertEqual(app.tabs.select(), str(app.notes_tab))
                self.assertEqual(app.reader_paper.current(), 1)
                app.reader_folder.current(0)
                app.reader_folder_selected()
                app.refresh()
                settle_library(app, root)
                root.update()
                self.assertTrue(app.all_projects.get())
                self.assertEqual(app.reader_folder.current(), 0)
                app.reader_folder.current(1)
                app.reader_folder_selected()
                app.reader_paper.current(0)
                app.reader_paper_selected()
                root.update()
                self.assertFalse(app.papers.selection())
                self.assertFalse(app.all_projects.get())
                app.search_fields = ["ocr"]
                app.search.set("control experiment")
                app.refresh_notes()
                settle_notes(app, root)
                self.assertEqual(app.entries, [])
                app.search_fields = ["quote"]
                app.search.set("stable OR")
                app.refresh_notes()
                settle_notes(app, root)
                self.assertTrue(app.search_error.get())
                self.assertEqual(app.entries, [])
                self.assertTrue(app.export_button.instate(["disabled"]))
                app.search.set("stable OR nonexistent")
                app.refresh_notes()
                settle_notes(app, root)
                self.assertFalse(app.search_error.get())
                self.assertEqual(len(app.entries), 1)
                self.assertFalse(app.export_button.instate(["disabled"]))
                app.entries[0]["label"] = "233"
                with patch.object(app.pdf_worker, "materialize", return_value=app.entries[0]["path"]), \
                        patch("note_organizer.ui.open_pdf_page", return_value=("Opened", True)) as opener:
                    app.open_source()
                    app.pdf_open_future.result(timeout=3)
                    app.poll()
                    opener.assert_called_once_with(Path(app.entries[0]["path"]), app.entries[0]["page"], app.viewer_preference)
                app.toggle_activity()
                self.assertEqual(app.activity_frame.winfo_manager(), "pack")
                app.toggle_activity()
                self.assertFalse(app.activity_frame.winfo_manager())
                self.assertFalse(errors)
        finally:
            for pending in root.tk.call("after", "info"):
                root.after_cancel(pending)
            root.destroy()

    def test_failed_first_scan_is_visible_and_theme_initializes(self):
        self.store.scan_result(self.pid, self.folder / "broken.pdf", "Invalid PDF")
        root = tk.Tk()
        root.withdraw()
        errors = []
        root.report_callback_exception = lambda *args: errors.append(args)
        try:
            with patch.object(Service, "start"):
                app = App(root, self.store, self.folder)
                settle_library(app, root)
                root.update()
                self.assertTrue(app.failed_papers)
                iid = next(iter(app.failed_papers))
                app.papers.selection_set(iid)
                root.update()
                app.refresh_notes()
                settle_notes(app, root)
                self.assertIn("Invalid PDF", app.preview_text.get("1.0", "end"))
                self.assertEqual(app.style.lookup("TButton", "foreground"), "#DCE4EE")
                self.assertEqual(app.preview_text.cget("background"), "#18191a")
                self.assertFalse(errors)
        finally:
            for pending in root.tk.call("after", "info"):
                root.after_cancel(pending)
            root.destroy()


if __name__ == "__main__":
    unittest.main()
