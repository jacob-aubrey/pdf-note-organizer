import base64
from io import BytesIO
import json
from pathlib import Path
import shutil
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

import fitz
from docx import Document
from PIL import Image

from note_organizer.extraction import extract, group_graphics, ocr_image
from note_organizer.exporting import export_notebook
from note_organizer.service import Service
from note_organizer.storage import Store, DEFAULT_SETTINGS, canonical, digest_file


def sample_pdf(path, rotate=0):
    with fitz.open() as doc:
        p = doc.new_page(width=612, height=792)
        p.insert_text((60, 65), "A study of catalyst stability", fontsize=18)
        p.insert_text((60, 105), "The catalyst remained stable after heating to 300 degrees.", fontsize=12)
        p.insert_text((60, 125), "Compare the marked result with the control experiment.", fontsize=12)
        a = p.add_highlight_annot(p.search_for("remained stable"))
        a.set_info(content="Check this claim")
        a.update()
        reply = p.add_text_annot((350, 110), "Compare replicates")
        reply.set_irt_xref(a.xref)
        reply.update()
        # Four separated letter-like annotations, all one handwritten phrase.
        for x in (65, 82, 99, 116):
            ink = p.add_ink_annot([[(x, 168), (x+5, 153), (x+10, 168)]])
            ink.set_colors(stroke=(1, 0, 0))
            ink.set_border(width=1.5)
            ink.update()
        arrow = p.add_line_annot((130, 157), (240, 108))
        arrow.set_colors(stroke=(1, 0, 0))
        arrow.set_line_ends(0, 4)
        arrow.update()
        # A distant annotation must remain separate.
        ink = p.add_ink_annot([[(420, 540), (440, 555), (460, 535)]])
        ink.update()
        p.insert_text((385, 585), "Figure 2 control response", fontsize=11)
        p.draw_rect(fitz.Rect(380, 500, 520, 565), color=(.2, .2, .2))
        p.set_rotation(rotate)
        doc.set_metadata({"title": "Catalyst stability", "author": "Example Researcher"})
        doc.save(path)


class NotebookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.pdf = self.folder / "paper.pdf"
        sample_pdf(self.pdf)
        self.store = Store(self.folder / "library.sqlite3")
        self.project_id = self.store.add_project(self.folder)
        self.service = Service(self.store)

    def tearDown(self):
        self.service.stop()
        self.tmp.cleanup()

    def scan(self, force=False):
        self.service.scan(self.store.project(self.project_id), self.pdf, force)
        return self.store.entries(self.project_id)

    def test_letters_and_arrow_are_one_group_with_original_text(self):
        _, _, _, entries = extract(self.pdf, DEFAULT_SETTINGS)
        drawings = [e for e in entries if e["kind"] == "drawing"]
        self.assertEqual(len(drawings), 2)
        self.assertIn("catalyst remained stable", drawings[0]["context"])
        self.assertGreater(len(drawings[0]["key"].split("+")), 4)
        with Image.open(BytesIO(base64.b64decode(drawings[0]["image"]))) as image:
            self.assertGreater(image.width, 300)

    def test_comments_and_reply_only_once(self):
        entries = self.scan()
        highlights = [e for e in entries if e["kind"] == "highlight"]
        self.assertEqual(highlights[0]["quote"], "remained stable")
        self.assertEqual(highlights[0]["comment"], "Check this claim\nCompare replicates")
        self.assertFalse(any(e["kind"] == "text" for e in entries))

    def test_repeat_save_does_not_duplicate(self):
        first = self.scan()
        with fitz.open(self.pdf) as doc:
            doc.set_metadata({**doc.metadata, "subject": "Saved again"})
            doc.saveIncr()
        second = self.scan()
        self.assertEqual(len(first), len(second))
        self.assertEqual(len(self.store.documents()), 1)

    def test_edit_and_delete_reconciled(self):
        self.scan()
        with fitz.open(self.pdf) as doc:
            p = doc[0]
            a = next(p.annots())
            a.set_info(content="Revised interpretation")
            a.update()
            doc.saveIncr()
        entries = self.scan()
        highlight = next(e for e in entries if e["kind"] == "highlight")
        self.assertIn("Revised interpretation", highlight["comment"])
        with fitz.open(self.pdf) as doc:
            p = doc[0]
            while p.first_annot:
                p.delete_annot(p.first_annot)
            doc.saveIncr()
        self.assertEqual(self.scan(), [])
        doc_id = self.store.documents()[0]["id"]
        actions = [h["action"] for h in self.store.history(doc_id)]
        self.assertIn("updated", actions)
        self.assertIn("removed", actions)

    def test_same_name_different_subfolders(self):
        self.scan()
        other = self.folder / "sub" / "paper.pdf"
        other.parent.mkdir()
        shutil.copy2(self.pdf, other)
        self.service.scan(self.store.project(self.project_id), other)
        self.assertEqual(len(self.store.documents()), 2)
        self.assertEqual(len(self.store.entries()), 6)

    def test_overlapping_projects_independent(self):
        self.scan()
        parent_id = self.store.add_project(self.folder.parent)
        self.service.scan(self.store.project(parent_id), self.pdf)
        self.assertEqual(len(self.store.entries(parent_id)), 3)
        self.assertEqual(len(self.store.entries(self.project_id)), 3)

    def test_rename_keeps_document_identity(self):
        self.scan()
        before = self.store.documents()[0]["id"]
        new = self.pdf.with_name("renamed.pdf")
        self.pdf.rename(new)
        self.service.scan(self.store.project(self.project_id), new)
        self.assertEqual(len(self.store.documents()), 1)
        self.assertEqual(self.store.documents()[0]["id"], before)
        self.assertEqual(self.store.entries()[0]["path"], canonical(new))

    def test_failed_scan_retains_last_good_notes(self):
        self.scan()
        self.pdf.write_bytes(b"broken")
        with self.assertRaises(Exception):
            self.scan()
        self.assertEqual(len(self.store.entries()), 3)

    def test_failed_export_stays_pending_and_recovers(self):
        self.scan()
        with self.assertLogs(level="ERROR"):
            with patch("note_organizer.exporting.os.replace", side_effect=PermissionError("Word holds file open")):
                self.service.export_pending()
        p = self.store.project(self.project_id)
        self.assertTrue(p["dirty"])
        self.assertIn("Export pending", p["error"])
        self.service.export_pending()
        p = self.store.project(self.project_id)
        self.assertFalse(p["dirty"])
        self.assertTrue(Path(p["output"]).exists())

    def test_rebuild_without_pdf_and_missing_source_retained(self):
        self.scan()
        self.pdf.unlink()
        self.store.mark_missing(self.project_id, set())
        self.service.export_pending()
        p = self.store.project(self.project_id)
        text = "\n".join(p.text for p in Document(p["output"]).paragraphs)
        self.assertIn("Source unavailable", text)
        self.assertIn("remained stable", text)

    def test_external_word_edits_are_preserved(self):
        self.scan()
        self.service.export_pending()
        p = self.store.project(self.project_id)
        doc = Document(p["output"])
        doc.add_paragraph("My irreplaceable research thought")
        doc.save(p["output"])
        self.store.mark_dirty(self.project_id)
        self.service.export_pending()
        backups = list((self.folder / "Notebook backups").glob("*.docx"))
        self.assertEqual(len(backups), 1)
        text = "\n".join(p.text for p in Document(backups[0]).paragraphs)
        self.assertIn("My irreplaceable research thought", text)

    def test_docx_has_images_and_source_links(self):
        self.scan()
        self.service.export_pending()
        p = self.store.project(self.project_id)
        with zipfile.ZipFile(p["output"]) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.startswith("word/media/")]), 2)
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertIn("#page=1", relationships)

    def test_whole_page_preserves_all_marks_together(self):
        _, _, _, entries = extract(self.pdf, {**DEFAULT_SETTINGS, "grouping": "Whole page"})
        drawings = [e for e in entries if e["kind"] == "drawing"]
        self.assertEqual(len(drawings), 1)
        self.assertEqual(drawings[0]["crop"], [0, 0, 612, 792])
        self.assertIn("Figure 2", drawings[0]["context"])

    def test_rotated_page_context_not_blank(self):
        rotated = self.folder / "rotated.pdf"
        sample_pdf(rotated, rotate=90)
        _, _, _, entries = extract(rotated, DEFAULT_SETTINGS)
        drawing = next(e for e in entries if e["kind"] == "drawing")
        with Image.open(BytesIO(base64.b64decode(drawing["image"]))) as image:
            self.assertLess(min(image.convert("L").getextrema()), 100)
        self.assertIn("catalyst remained stable", drawing["context"])

    def test_quiet_file_is_due_even_after_long_queue_wait(self):
        self.service.enqueue(self.pdf)
        self.assertEqual(self.service.due(time.monotonic()), [])
        ready = self.service.due(time.monotonic()+10)
        self.assertEqual(ready, [(canonical(self.pdf), 0)])

    def test_events_coalesce_and_reschedule(self):
        with patch("note_organizer.service.time.monotonic", return_value=10):
            self.service.enqueue(self.pdf)
        with patch("note_organizer.service.time.monotonic", return_value=11):
            self.service.enqueue(self.pdf)
        self.assertEqual(self.service.due(12), [])
        self.assertEqual(len(self.service.due(13)), 1)

    def test_missing_ocr_does_not_lose_images(self):
        with patch("note_organizer.extraction.shutil.which", return_value=None):
            _, _, meta, entries = extract(self.pdf, {**DEFAULT_SETTINGS, "ocr": True})
        self.assertTrue(meta["warnings"])
        self.assertTrue(all(e["image"] for e in entries))

    def test_legacy_import_read_only_and_idempotent(self):
        legacy = self.folder / "legacy.db"
        with sqlite3.connect(legacy) as db:
            db.execute("CREATE TABLE watched_folders(root_path TEXT,notes_path TEXT)")
            db.execute("INSERT INTO watched_folders VALUES(?,?)", (str(self.folder), "old notebook.docx"))
        db.close()
        before = digest_file(legacy)
        self.store.import_legacy([legacy])
        self.store.import_legacy([legacy])
        self.assertEqual(len(self.store.projects()), 1)
        self.assertEqual(digest_file(legacy), before)

    def test_removing_project_cascades_library_but_leaves_files(self):
        self.scan()
        self.service.export_pending()
        output = Path(self.store.project(self.project_id)["output"])
        self.store.remove_project(self.project_id)
        self.assertEqual(self.store.entries(), [])
        self.assertEqual(self.store.documents(), [])
        self.assertTrue(output.exists())
        self.assertTrue(self.pdf.exists())

    def test_search_in_comments_colors_and_metadata(self):
        self.scan()
        self.assertEqual(len(self.store.entries(query="replicates")), 1)
        self.assertEqual(len(self.store.entries(query="Example Researcher")), 3)
        self.assertGreater(len(self.store.entries(query="Red")), 0)

    def test_live_watcher_catches_saved_change_and_shutdown(self):
        self.service.debounce = .1
        self.service.start()

        def wait_for(predicate):
            deadline = time.monotonic() + 12
            while time.monotonic() < deadline:
                if predicate():
                    return
                time.sleep(.03)
            self.fail("Background service did not finish within 12 seconds")

        try:
            wait_for(lambda: self.store.entries() and not self.store.project(self.project_id)["dirty"])
            with fitz.open(self.pdf) as doc:
                p = doc[0]
                a = next(p.annots())
                a.set_info(content="Live watcher caught this")
                a.update()
                doc.saveIncr()
            wait_for(lambda: any("Live watcher caught this" in e["comment"] for e in self.store.entries()))
            wait_for(lambda: not self.store.project(self.project_id)["dirty"])
            output = self.store.project(self.project_id)["output"]
            self.assertIn("Live watcher caught this", " ".join(p.text for p in Document(output).paragraphs))
        finally:
            self.service.stop()
            self.service.thread.join(timeout=10)
        self.assertFalse(self.service.thread.is_alive())

    def test_startup_catches_changes_made_while_closed(self):
        self.scan()
        with fitz.open(self.pdf) as doc:
            p = doc[0]
            a = p.add_text_annot((100, 400), "Added while app was closed")
            a.update()
            doc.saveIncr()
        self.service.sync()
        self.assertTrue(self.store.entries(query="Added while app was closed"))

    def test_scanned_highlight_has_visual_fallback(self):
        scanned = self.folder / "scanned.pdf"
        with fitz.open() as doc:
            p = doc.new_page()
            image = Image.new("RGB", (300, 80), "white")
            from PIL import ImageDraw
            ImageDraw.Draw(image).text((10, 25), "Image-only source text", fill="black")
            buffer = BytesIO()
            image.save(buffer, format="PNG")
            p.insert_image(fitz.Rect(40, 50, 340, 130), stream=buffer.getvalue())
            a = p.add_highlight_annot(fitz.Rect(45, 70, 230, 100))
            a.update()
            doc.save(scanned)
        _, _, metadata, entries = extract(scanned, DEFAULT_SETTINGS)
        self.assertEqual(entries[0]["quote"], "")
        self.assertTrue(entries[0]["image"])
        self.assertTrue(metadata["warnings"])

    def test_column_highlight_does_not_collect_neighbor_column(self):
        path = self.folder / "columns.pdf"
        with fitz.open() as doc:
            p = doc.new_page()
            p.insert_text((40, 90), "Selected words in left column")
            p.insert_text((330, 90), "Unrelated right column")
            p.add_highlight_annot(p.search_for("Selected words"))
            doc.save(path)
        _, _, _, entries = extract(path, DEFAULT_SETTINGS)
        self.assertEqual(entries[0]["quote"], "Selected words")

    def test_failed_replace_keeps_existing_notebook_intact(self):
        self.scan()
        self.service.export_pending()
        target = self.store.project(self.project_id)["output"]
        original = digest_file(target)
        self.store.mark_dirty(self.project_id)
        with self.assertLogs(level="ERROR"):
            with patch("note_organizer.exporting.os.replace", side_effect=PermissionError("Locked")):
                self.service.export_pending()
        self.assertEqual(digest_file(target), original)
        self.assertEqual(list(self.folder.glob(".pdf-notes-*.docx")), [])

    def test_context_includes_complete_intersecting_text_lines(self):
        _, _, _, entries = extract(self.pdf, DEFAULT_SETTINGS)
        e = next(e for e in entries if e["kind"] == "drawing")
        self.assertIn("heating to 300 degrees.", e["context"])
        self.assertIn("control experiment.", e["context"])


if __name__ == "__main__":
    unittest.main()
