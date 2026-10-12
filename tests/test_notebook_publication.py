"""Notebook publication races use only disposable synthetic Word documents."""
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from note_organizer import exporting, remarkable_cloud
from note_organizer.storage import digest_file


TEMP_ROOT = Path(__file__).resolve().parents[1] / "tmp" / "combined-tests"


class NotebookPublicationTests(unittest.TestCase):
    def setUp(self):
        TEMP_ROOT.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=TEMP_ROOT)
        self.root = Path(self.temp.name)
        self.path = self.root / "space \u03bb \U0001f680.docx"
        self.editor = self.root / "editor.docx"
        self.write_document(self.editor, "Concurrent user writing")
        self.editor_bytes = self.editor.read_bytes()

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def write_document(path, text):
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    def publish_with_concurrent_creation(self, expected_hash=""):
        """Inject after the old final hash or after the new guarded move."""
        injected = False
        old_replace = os.replace
        old_publish = remarkable_cloud._publish_absent

        def race(function, temporary, destination):
            nonlocal injected
            if Path(destination) == self.path and not injected:
                injected = True
                self.path.write_bytes(self.editor_bytes)
            return function(temporary, destination)

        error = None
        with patch.object(exporting.os, "replace", side_effect=lambda a, b: race(old_replace, a, b)):
            with patch.object(remarkable_cloud, "_publish_absent", side_effect=lambda a, b: race(old_publish, a, b)):
                try:
                    exporting.export_notebook(self.path, "Generated notebook", [], expected_hash)
                except (OSError, remarkable_cloud.CloudError) as failure:
                    error = failure
        self.assertTrue(injected)
        return error

    def test_missing_target_created_at_publication_is_preserved(self):
        error = self.publish_with_concurrent_creation()
        self.assertEqual(self.path.read_bytes(), self.editor_bytes)
        self.assertIsInstance(error, OSError)
        self.assertIn("Notebook", str(error))
        self.assertEqual(list(self.root.glob("*.preserved")), [])
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_managed_target_created_after_final_hash_retains_both_versions(self):
        self.write_document(self.path, "Previous generated notebook")
        original = self.path.read_bytes()
        error = self.publish_with_concurrent_creation(digest_file(self.path))
        self.assertEqual(self.path.read_bytes(), self.editor_bytes)
        preserved = list(self.root.glob("*.preserved"))
        self.assertEqual(len(preserved), 1)
        self.assertEqual(preserved[0].read_bytes(), original)
        self.assertIsInstance(error, OSError)
        self.assertIn("notebook", str(error).lower())
        self.assertIn(str(preserved[0]), str(error))
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_edit_after_final_hash_before_guarded_open_is_preserved(self):
        self.write_document(self.path, "Previous generated notebook")
        expected = digest_file(self.path)
        real = exporting.publish_staged

        def edit_before_guard(temporary, path, **options):
            self.path.write_bytes(self.editor_bytes)
            return real(temporary, path, **options)

        with patch.object(exporting, "publish_staged", side_effect=edit_before_guard):
            with self.assertRaisesRegex(OSError, "Notebook.*notebook changed"):
                exporting.export_notebook(self.path, "Generated notebook", [], expected)
        self.assertEqual(self.path.read_bytes(), self.editor_bytes)
        self.assertEqual(list(self.root.glob("*.preserved")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_existing_native_writable_handle_prevents_notebook_replacement(self):
        self.write_document(self.path, "Previous generated notebook")
        original = self.path.read_bytes()
        expected = digest_file(self.path)
        with self.path.open("r+b"):
            with self.assertRaisesRegex(OSError, "Notebook.*open for editing"):
                exporting.export_notebook(self.path, "Generated notebook", [], expected)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.root.glob("*.preserved")), [])
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_preexisting_edited_notebook_keeps_revision_backup(self):
        self.path.write_bytes(self.editor_bytes)
        edited_hash = digest_file(self.path)
        saved_hash, backup = exporting.export_notebook(self.path, "Generated notebook", [], "old managed hash")
        self.assertEqual(saved_hash, digest_file(self.path))
        self.assertEqual(Path(backup), self.root / "Notebook backups" / f"{self.path.stem}.{edited_hash}.docx")
        self.assertEqual(Path(backup).read_bytes(), self.editor_bytes)
        self.assertNotEqual(self.path.read_bytes(), self.editor_bytes)
        self.assertEqual(list(self.root.glob("*.preserved")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_successful_cleanup_failure_returns_exact_previous_copy(self):
        self.write_document(self.path, "Previous generated notebook")
        original = self.path.read_bytes()
        expected = digest_file(self.path)
        with patch.object(remarkable_cloud, "_delete_guarded", side_effect=PermissionError("Synthetic cleanup lock")):
            saved_hash, backup = exporting.export_notebook(self.path, "Generated notebook", [], expected)
        self.assertEqual(saved_hash, digest_file(self.path))
        self.assertNotEqual(self.path.read_bytes(), original)
        self.assertEqual(Path(backup).read_bytes(), original)
        self.assertEqual(list(self.root.glob("*.preserved")), [Path(backup)])
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_cleanup_failure_prefers_preserved_location_and_keeps_edited_backup(self):
        self.path.write_bytes(self.editor_bytes)
        with patch.object(remarkable_cloud, "_delete_guarded", side_effect=PermissionError("Synthetic cleanup lock")):
            saved_hash, backup = exporting.export_notebook(self.path, "Generated notebook", [], "old managed hash")
        self.assertEqual(saved_hash, digest_file(self.path))
        self.assertTrue(backup.endswith(".docx.preserved") or backup.endswith(".preserved"))
        self.assertEqual(Path(backup).read_bytes(), self.editor_bytes)
        revisions = list((self.root / "Notebook backups").glob("*.docx"))
        self.assertEqual(len(revisions), 1)
        self.assertEqual(revisions[0].read_bytes(), self.editor_bytes)
        self.assertEqual(list(self.root.glob("*.preserved")), [Path(backup)])

    def test_staged_notebook_is_fsynced_before_publication(self):
        real = remarkable_cloud._publish_absent
        with patch.object(remarkable_cloud.os, "fsync", wraps=os.fsync) as flush:
            def verify_flush(temporary, path):
                self.assertEqual(flush.call_count, 1)
                self.assertIsInstance(Document(temporary), type(Document()))
                return real(temporary, path)
            with patch.object(remarkable_cloud, "_publish_absent", side_effect=verify_flush):
                saved_hash, backup = exporting.export_notebook(self.path, "Generated notebook", [])
        self.assertEqual(saved_hash, digest_file(self.path))
        self.assertEqual(backup, "")

    def test_failed_staged_flush_leaves_original_untouched(self):
        self.write_document(self.path, "Previous generated notebook")
        original = self.path.read_bytes()
        expected = digest_file(self.path)
        with patch.object(remarkable_cloud.os, "fsync", side_effect=OSError("Synthetic disk flush failure")):
            with self.assertRaisesRegex(OSError, "Notebook.*disk flush failure"):
                exporting.export_notebook(self.path, "Generated notebook", [], expected)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.root.glob("*.preserved")), [])
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_publication_failure_restores_original_and_reports_notebook_context(self):
        self.write_document(self.path, "Previous generated notebook")
        original = self.path.read_bytes()
        expected = digest_file(self.path)
        with patch.object(remarkable_cloud, "_publish_absent", side_effect=OSError("Synthetic publication failure")):
            with self.assertRaisesRegex(OSError, "Notebook.*publication failure"):
                exporting.export_notebook(self.path, "Generated notebook", [], expected)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.root.glob("*.preserved")), [])

    def test_interrupted_notebook_replacement_reports_retained_path(self):
        preserved = self.path.with_name(self.path.name + ".synthetic.preserved")
        preserved.write_bytes(self.editor_bytes)
        with self.assertRaisesRegex(OSError, "Notebook.*interrupted replacement") as failure:
            exporting.export_notebook(self.path, "Generated notebook", [])
        self.assertIn(str(preserved), str(failure.exception))
        self.assertFalse(self.path.exists())
        self.assertEqual(preserved.read_bytes(), self.editor_bytes)
        self.assertEqual(list(self.root.glob(".pdf-notes-*")), [])

    @unittest.skipUnless(os.name == "nt", "Mandatory Windows sharing guards")
    def test_edit_after_publication_is_backed_up_on_next_export(self):
        real = exporting.publish_staged
        generated = None

        def edit_after_publication(temporary, path, **options):
            nonlocal generated
            generated = digest_file(temporary)
            retained = real(temporary, path, **options)
            self.path.write_bytes(self.editor_bytes)
            return retained

        with patch.object(exporting, "publish_staged", side_effect=edit_after_publication):
            saved_hash, backup = exporting.export_notebook(self.path, "Generated notebook", [])
        self.assertEqual(saved_hash, generated)
        self.assertNotEqual(saved_hash, digest_file(self.path))
        self.assertEqual(backup, "")
        _, backup = exporting.export_notebook(self.path, "Generated notebook", [], saved_hash)
        self.assertEqual(Path(backup).read_bytes(), self.editor_bytes)


if __name__ == "__main__":
    unittest.main()
