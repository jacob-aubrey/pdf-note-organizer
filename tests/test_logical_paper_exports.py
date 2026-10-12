"""Logical paper exports preserve source-specific pages, links, and previews."""
import base64
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from PIL import Image

from note_organizer.exporting import build_notebook, file_uri, ordered_entries
from note_organizer.paper_registry import Registry
from note_organizer.storage import Store, canonical


class LogicalPaperExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.original = self.folder / "A Original.pdf"
        self.tablet = self.folder / "Z Tablet copy.pdf"
        self.other = self.folder / "B Different paper.pdf"

    def tearDown(self):
        self.temp.cleanup()

    def entry(self, path, text, **extra):
        return {"path": str(path), "key": "0:stable-mark", "document_id": path.stem,
                "source_digest": "source-" + path.stem, "metadata": {"title": path.stem},
                "kind": "highlight", "page": 1, "page_count": 1, "rect": [10, 10, 30, 30],
                "quote": text, "comment": "", "image": "", "ocr": "", **extra}

    def linked_entries(self):
        return [self.entry(self.tablet, "Tablet notes", logical_paper_id="one-paper", logical_paper_title="Shared paper title"),
                self.entry(self.other, "Separate paper notes"),
                self.entry(self.original, "Original Acrobat notes", logical_paper_id="one-paper", logical_paper_title="Shared paper title")]

    @staticmethod
    def serialized(document):
        stream = BytesIO()
        document.save(stream)
        return zipfile.ZipFile(BytesIO(stream.getvalue()))

    def test_each_layout_emits_one_heading_for_linked_sources_and_keeps_other_paper(self):
        for layout in ("Reading", "Detailed"):
            with self.subTest(layout=layout):
                document = build_notebook("Folder notes", self.linked_entries(), layout=layout)
                headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
                self.assertEqual(headings, ["Shared paper title", "B Different paper"])
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                for excerpt in ("Tablet notes", "Original Acrobat notes", "Separate paper notes"):
                    self.assertIn(excerpt, text)
                with self.serialized(document) as archive:
                    relationships = archive.read("word/_rels/document.xml.rels").decode()
                    for path in (self.original, self.tablet, self.other):
                        self.assertIn(file_uri(path, 1), relationships)

    def test_reading_pages_remain_separate_for_original_and_tablet_versions(self):
        entries = [entry for entry in self.linked_entries() if entry.get("logical_paper_id")]
        document = build_notebook("Folder notes", entries)
        pages = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 2"]
        self.assertEqual(len(pages), 2)
        xml = "\n".join(paragraph._p.xml for paragraph in pages)
        self.assertIn("A Original.pdf", xml)
        self.assertIn("Z Tablet copy.pdf", xml)
        with self.serialized(document) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertIn(file_uri(self.original, 1), relationships)
            self.assertIn(file_uri(self.tablet, 1), relationships)

    def test_page_revisions_remain_distinct_even_when_the_path_and_page_match(self):
        first = self.entry(self.original, "First version", logical_paper_id="paper", source_digest="revision-one")
        second = {**first, "quote": "Second version", "source_digest": "revision-two"}
        document = build_notebook("Folder notes", [first, second])
        self.assertEqual(len([p for p in document.paragraphs if p.style.name == "Heading 2"]), 2)

    def test_verified_combined_notes_link_only_durable_original_and_share_page_heading(self):
        entries = [self.entry(self.original, "PC comment", logical_paper_id="paper",
                              source_link_path=str(self.original), source_link_page=1),
                   self.entry(self.tablet, "Tablet comment", logical_paper_id="paper",
                              source_link_path=str(self.original), source_link_page=1)]
        original_entries = deepcopy(entries)
        for layout in ("Reading", "Detailed"):
            with self.subTest(layout=layout):
                document = build_notebook("Combined notes", entries, layout=layout)
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn("PC comment", text)
                self.assertIn("Tablet comment", text)
                self.assertIn("For verified papers, PDF links open the saved original paper", text)
                if layout == "Reading":
                    self.assertEqual(len([p for p in document.paragraphs if p.style.name == "Heading 2"]), 1)
                with self.serialized(document) as archive:
                    relationships = archive.read("word/_rels/document.xml.rels").decode()
                    self.assertIn(file_uri(self.original, 1), relationships)
                    self.assertNotIn(file_uri(self.tablet, 1), relationships)
        self.assertEqual(entries, original_entries)

    def test_unverified_page_has_no_link_to_disposable_or_wrong_original_page(self):
        entries = [self.entry(self.tablet, "Added tablet page", logical_paper_id="paper",
                              source_link_path=None, source_link_page=None)]
        for layout in ("Reading", "Detailed"):
            with self.subTest(layout=layout):
                document = build_notebook("Notes needing review", entries, layout=layout)
                self.assertIn("source page needs review", "\n".join(p.text for p in document.paragraphs))
                with self.serialized(document) as archive:
                    relationships = archive.read("word/_rels/document.xml.rels").decode()
                    self.assertNotIn(file_uri(self.tablet, 1), relationships)
                    self.assertNotIn(file_uri(self.original, 1), relationships)

    def test_store_pc_and_managed_tablet_notes_share_one_durable_page_heading(self):
        store = Store(self.folder / "data" / "notes.sqlite3")
        project = store.add_project(self.folder)
        self.original.write_bytes(b"Original PDF bytes")
        anchor = self.folder / "data" / "remarkable" / "documents" / "cloud-paper.pdf"
        registry = Registry(store)
        original = registry.register_local(self.original)
        registry.link_local(original["paper_id"], anchor, digest="tablet-digest",
                            metadata={"role": "tablet_render", "cloud_id": "cloud-paper"})
        store.save_scan(project, self.original, "pc-digest", "signature", "{}", {},
                        [self.entry(self.original, "PC comment")])
        store.save_scan(project, anchor, "tablet-digest", "signature", "{}",
                        {"managed_cloud_id": "cloud-paper", "primary_path": canonical(self.original)},
                        [self.entry(anchor, "Tablet comment")])
        document_ids = {row["id"] for row in store.documents(project)}
        histories = {doc_id: store.history(doc_id) for doc_id in document_ids}

        entries = store.entries(project)
        self.assertEqual({entry["document_id"] for entry in entries}, document_ids)
        self.assertTrue(all(entry["source_link_path"] == canonical(self.original) for entry in entries))
        self.assertTrue(all(entry["source_link_page"] == 1 for entry in entries))
        document = build_notebook("Combined saved notes", entries)
        self.assertEqual(len([p for p in document.paragraphs if p.style.name == "Heading 2"]), 1)
        with self.serialized(document) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertIn(file_uri(self.original, 1), relationships)
            self.assertNotIn(file_uri(anchor, 1), relationships)
        self.assertEqual({doc_id: store.history(doc_id) for doc_id in document_ids}, histories)

    def test_store_page_normalization_keeps_legacy_variants_and_other_projects_separate(self):
        store = Store(self.folder / "data" / "notes.sqlite3")
        project = store.add_project(self.folder)
        other_project = store.add_project(self.folder / "Other project")
        self.original.write_bytes(b"Original PDF bytes")
        self.tablet.write_bytes(b"Independent legacy reading PDF")
        anchor = self.folder / "data" / "remarkable" / "documents" / "cloud-paper.pdf"
        registry = Registry(store)
        original = registry.register_local(self.original)
        for path in (anchor, self.tablet):
            registry.link_local(original["paper_id"], path, digest="tablet-digest",
                                metadata={"role": "tablet_render", "cloud_id": "cloud-paper"})
        for membership in (project, other_project):
            store.save_scan(membership, self.original, "pc-digest", "signature", "{}", {},
                            [self.entry(self.original, "PC comment")])
        store.save_scan(project, self.tablet, "legacy-digest", "signature", "{}", {},
                        [self.entry(self.tablet, "Legacy comment")])
        store.save_scan(project, anchor, "tablet-digest", "signature", "{}",
                        {"managed_cloud_id": "cloud-paper", "primary_path": canonical(self.original)},
                        [self.entry(anchor, "Tablet comment")])

        entries = store.entries(project)
        legacy = next(entry for entry in entries if entry["quote"] == "Legacy comment")
        self.assertNotIn("source_link_path", legacy)
        self.assertNotIn("source_link_page", legacy)
        self.assertNotIn("source_link_path", store.entries(other_project)[0])
        document = build_notebook("Mixed source notes", entries)
        self.assertEqual(len([p for p in document.paragraphs if p.style.name == "Heading 2"]), 2)
        guidance = next(p.text for p in document.paragraphs if p.text.startswith("For verified papers,"))
        self.assertIn("Older tablet copies keep links to their retained reading PDFs until a successful refresh.", guidance)
        with self.serialized(document) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertIn(file_uri(self.original, 1), relationships)
            self.assertIn(file_uri(self.tablet, 1), relationships)
            self.assertNotIn(file_uri(anchor, 1), relationships)

    def test_mapped_page_link_uses_verified_original_page(self):
        entry = self.entry(self.tablet, "Mapped tablet page", page=3, logical_paper_id="paper",
                           source_link_path=str(self.original), source_link_page=2)
        for layout in ("Reading", "Detailed"):
            with self.subTest(layout=layout):
                with self.serialized(build_notebook("Notes", [entry], layout=layout)) as archive:
                    relationships = archive.read("word/_rels/document.xml.rels").decode()
                    self.assertIn(file_uri(self.original, 2), relationships)
                    self.assertNotIn(file_uri(self.tablet, 3), relationships)

    def test_drawing_previews_from_different_variants_never_merge(self):
        entries = []
        for path, color in ((self.original, "red"), (self.tablet, "blue")):
            stream = BytesIO()
            Image.new("RGB", (20, 20), color).save(stream, format="PNG")
            entries.append(self.entry(path, "", kind="drawing", crop=[10, 10, 30, 30],
                image=base64.b64encode(stream.getvalue()).decode(), logical_paper_id="paper", logical_paper_title="One paper"))
        before = deepcopy(entries)
        for layout in ("Reading", "Detailed"):
            with self.subTest(layout=layout):
                document = build_notebook("Folder notes", entries, layout=layout)
                with self.serialized(document) as archive:
                    self.assertEqual(len([name for name in archive.namelist() if name.startswith("word/media/")]), 2)
                    relationships = archive.read("word/_rels/document.xml.rels").decode()
                    self.assertIn(file_uri(self.original, 1), relationships)
                    self.assertIn(file_uri(self.tablet, 1), relationships)
        self.assertEqual(entries, before)

    def test_unregistered_export_sort_and_headings_keep_existing_behavior(self):
        entries = [self.entry(self.tablet, "Z"), self.entry(self.original, "A"), self.entry(self.other, "B")]
        self.assertEqual([entry["path"] for entry in ordered_entries(entries)],
                         [str(self.original), str(self.other), str(self.tablet)])
        document = build_notebook("Folder notes", entries)
        self.assertEqual([p.text for p in document.paragraphs if p.style.name == "Heading 1"],
                         ["A Original", "B Different paper", "Z Tablet copy"])

    def test_confirmed_link_marks_only_affected_project_memberships_dirty(self):
        store = Store(self.folder / "data" / "notes.sqlite3")
        parent = store.add_project(self.folder)
        child = store.add_project(self.folder / "Child")
        unrelated = store.add_project(self.folder / "Unrelated")
        self.original.write_bytes(b"Retained original")
        self.other.write_bytes(b"Unrelated paper")
        for project, path in ((parent, self.original), (child, self.original), (unrelated, self.other)):
            store.save_scan(project, path, "digest", "signature", "{}", {}, [])
            store.export_result(project, "previous-export")
        registry = Registry(store)
        local = registry.register_local(self.original)
        cloud = registry.register_cloud("cloud-paper")
        registry.link_sources([local["id"], cloud["id"]])
        self.assertTrue(store.project(parent)["dirty"])
        self.assertTrue(store.project(child)["dirty"])
        self.assertFalse(store.project(unrelated)["dirty"])
        # Reconfirming the same relationship creates no extra export work.
        store.export_result(parent, "updated-export")
        store.export_result(child, "updated-export")
        registry.link_sources([local["id"], cloud["id"]])
        self.assertFalse(store.project(parent)["dirty"])
        self.assertFalse(store.project(child)["dirty"])

    def test_new_generated_variant_registration_dirties_only_related_notebook(self):
        store = Store(self.folder / "data" / "notes.sqlite3")
        related = store.add_project(self.folder / "Linked")
        unrelated = store.add_project(self.folder / "Unrelated")
        self.original.write_bytes(b"Original")
        self.tablet.write_bytes(b"Tablet PDF")
        for project, path in ((related, self.original), (related, self.tablet), (unrelated, self.other)):
            store.save_scan(project, path, "digest", "signature", "{}", {}, [])
        registry = Registry(store)
        original = registry.register_local(self.original)
        registry.link_cloud(original["paper_id"], "cloud-paper", digest="baseline")
        store.export_result(related, "current-notebook")
        store.export_result(unrelated, "unrelated-notebook")
        registry.link_local(original["paper_id"], self.tablet, metadata={"role": "tablet_render", "cloud_id": "cloud-paper"})
        self.assertTrue(store.project(related)["dirty"])
        self.assertFalse(store.project(unrelated)["dirty"])


if __name__ == "__main__":
    unittest.main()
