"""Offline Inbox integration tests; all documents and settings are synthetic."""
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from pypdf import PdfWriter

from note_organizer.inbox import DownloadNotReady, InboxController, file_signature
from note_organizer import remarkable_cloud
from note_organizer.storage import Store, canonical, timestamp
from note_organizer.sorting import core_logic
from tests.test_sorting_presentation_extraction import _write_pptx


class InboxTests(unittest.TestCase):
    def setUp(self):
        temporary_root = Path(__file__).resolve().parents[1] / "tmp"
        temporary_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="inbox-test-", dir=temporary_root)
        self.root = Path(self.temp.name)
        self.incoming = self.root / "ToSort"
        self.library = self.root / "Topics"
        self.data = self.root / "data"
        for folder in (self.incoming, self.library, self.data):
            folder.mkdir()
        self.store = Store(self.data / "library.sqlite")
        self.events = []
        self.inbox = InboxController(self.store, self.data, lambda event, payload: self.events.append((event, payload)))
        self.inbox.save_config({"watch_folder": str(self.incoming), "library_root": str(self.library),
                               "naming_mode": "Basic"})
        self.inbox.stable_seconds = 0

    def tearDown(self):
        self.inbox.stop()
        self.temp.cleanup()

    def pdf(self, name="incoming.pdf", folder=None):
        path = (folder or self.incoming) / name
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.add_metadata({"/Title": "Synthetic Research Paper", "/Author": "Jane Example"})
        writer.write(path)
        return path

    def ready(self, path):
        self.inbox.refresh()
        self.inbox.refresh()
        return self.inbox.prepare(path)

    def test_new_library_does_not_watch_or_inherit_external_settings(self):
        fresh = InboxController(Store(self.root / "fresh" / "library.sqlite"), self.root / "fresh")
        try:
            self.assertFalse(fresh.start())
            self.assertIsNone(fresh._thread)
            self.assertIsNone(fresh._executor)
            self.assertEqual(fresh.config()["watch_folder"], "")
            self.assertFalse(fresh.config()["enabled"])
        finally:
            fresh.stop()

    def test_construction_and_cached_config_do_not_access_locked_database(self):
        with patch.object(self.store, "connect", side_effect=AssertionError("Tk must not open the database")), \
             patch.object(self.store, "preference", side_effect=AssertionError("Tk must not read preferences")):
            controller = InboxController(self.store, self.data)
            try:
                self.assertFalse(controller.config()["enabled"])
                controller.set_config_snapshot({"watch_folder": str(self.incoming), "library_root": str(self.library),
                                                "enabled": False, "api_key_protected": "encrypted"})
                self.assertEqual(controller.config()["watch_folder"], str(self.incoming))
                self.assertTrue(controller.config()["api_key_present"])
                self.assertNotIn("encrypted", json.dumps(controller.config()))
                self.assertFalse(controller.start())
            finally:
                controller.stop()

    def test_scan_is_local_and_does_not_extract_or_request_metadata(self):
        self.pdf()
        with patch.object(core_logic, "get_basic_document_details", side_effect=AssertionError("not on scan")), \
             patch.object(core_logic, "get_document_details", side_effect=AssertionError("not on scan")):
            self.inbox.refresh()
            self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")
        self.assertFalse(any(event == "inbox_filed" for event, _ in self.events))

    def test_partial_pdf_waits_until_stable_eof(self):
        path = self.incoming / "download.pdf"
        path.write_bytes(b"%PDF-1.7\nstill downloading")
        self.inbox.refresh()
        self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "waiting")
        with self.assertRaises(DownloadNotReady):
            self.inbox.prepare(path)
        self.pdf(path.name)
        self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "waiting")
        self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")

    def test_stability_requires_elapsed_time_even_with_complete_pdf(self):
        self.inbox.stable_seconds = 2
        self.pdf()
        with patch("note_organizer.inbox.time.monotonic", side_effect=[10, 11, 12]):
            self.inbox.refresh()
            self.inbox.refresh()
            self.assertEqual(self.inbox.list_items()[0]["status"], "waiting")
            self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")

    def test_partial_incremental_pdf_does_not_accept_an_old_eof(self):
        path = self.pdf()
        path.write_bytes(path.read_bytes() + b"\n1 0 obj\n<< /Unfinished incremental update")
        self.inbox.refresh()
        self.inbox.refresh()
        self.assertEqual(self.inbox.list_items()[0]["status"], "waiting")
        with self.assertRaises(DownloadNotReady):
            self.inbox.prepare(path)

    def test_scan_ignores_office_lock_files_and_other_file_types(self):
        for name in ("~$open.docx", "source.rmdoc", "paper.pdf.crdownload", "download.part", "notes.txt"):
            (self.incoming / name).write_bytes(b"unprocessed")
        nested = self.incoming / "nested"
        nested.mkdir()
        self.pdf(folder=nested)
        self.assertEqual(self.inbox.refresh(), [])

    def test_preparing_status_recovers_after_restart(self):
        path = self.pdf()
        self.ready(path)
        with self.store.connect() as db:
            db.execute("UPDATE inbox_items SET status='preparing'")
        restarted = InboxController(self.store, self.data)
        try:
            self.assertEqual(restarted.list_items()[0]["status"], "waiting")
            restarted.stable_seconds = 0
            restarted.refresh()
            restarted.refresh()
            self.assertEqual(restarted.prepare(path)["source"], "Basic")
        finally:
            restarted.stop()

    def test_local_mode_never_calls_online_path(self):
        path = self.pdf()
        with patch.object(core_logic, "get_document_details", side_effect=AssertionError("must stay local")):
            details = self.ready(path)
        self.assertEqual(details["source"], "Basic")
        self.assertTrue(details["proposed_filename"].endswith(".pdf"))
        self.assertEqual(details["signature"], file_signature(path))

    def test_saved_gemini_key_is_not_decrypted_or_used_while_disabled(self):
        path = self.pdf()
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"):
            self.inbox.save_config({"api_key": "synthetic-own-key", "naming_mode": "Automatic"})
        fake = {"title": "Synthetic paper", "document_type": "journal_article", "year": "2026"}
        with patch("note_organizer.inbox._unprotect_for_current_windows_user", side_effect=AssertionError("disabled key")), \
             patch.object(core_logic, "get_document_details", return_value=fake) as read:
            self.ready(path)
        self.assertEqual(read.call_args.args[1], "")
        self.assertFalse(read.call_args.kwargs["allow_cloud_ai"])
        self.assertTrue(self.inbox.config()["api_key_present"])
        self.assertFalse(self.inbox.config()["gemini_enabled"])

    def test_explicit_gemini_enable_uses_current_windows_users_key(self):
        path = self.pdf()
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"), \
             patch("note_organizer.inbox._unprotect_for_current_windows_user", return_value="synthetic-own-key"):
            self.inbox.save_config({"api_key": "synthetic-own-key", "gemini_enabled": True, "naming_mode": "Automatic"})
            fake = {"title": "Synthetic paper", "document_type": "journal_article", "year": "2026"}
            with patch.object(core_logic, "get_document_details", return_value=fake) as read:
                self.ready(path)
        self.assertEqual(read.call_args.args[1], "synthetic-own-key")
        self.assertTrue(read.call_args.kwargs["allow_cloud_ai"])

    def test_gemini_enable_requires_an_available_key_for_this_user(self):
        with self.assertRaises(ValueError):
            self.inbox.save_config({"gemini_enabled": True})
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"):
            self.inbox.save_config({"api_key": "synthetic-own-key"})
        with patch("note_organizer.inbox._unprotect_for_current_windows_user", side_effect=OSError("other account")):
            with self.assertRaises(ValueError):
                self.inbox.save_config({"gemini_enabled": True})
        self.assertFalse(self.inbox.config()["gemini_enabled"])

    def test_cloud_consent_is_separate_for_each_office_format(self):
        docx = self.incoming / "report.docx"
        document = Document()
        document.add_paragraph("Synthetic report")
        document.save(docx)
        pptx = self.incoming / "slides.pptx"
        _write_pptx(pptx, ["Synthetic research presentation"])
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"), \
             patch("note_organizer.inbox._unprotect_for_current_windows_user", return_value="synthetic-key"):
            self.inbox.save_config({"naming_mode": "Automatic", "online_metadata_lookup_enabled": False,
                                   "gemini_enabled": True, "api_key": "synthetic-key"})
        fake = {"title": "Synthetic report", "document_type": "technical_report", "year": "2026"}
        with patch.object(core_logic, "get_document_details", return_value=fake) as read:
            self.ready(docx)
            self.assertFalse(read.call_args.kwargs["allow_cloud_ai"])
            self.assertFalse(read.call_args.kwargs["allow_online_metadata_lookup"])
            self.ready(pptx)
            self.assertFalse(read.call_args.kwargs["allow_cloud_ai"])
        with patch("note_organizer.inbox._unprotect_for_current_windows_user", return_value="synthetic-key"):
            self.inbox.save_config({"allow_cloud_ai_for_word_documents": True})
        with self.store.connect() as db:
            db.execute("UPDATE inbox_items SET status='ready',details='{}'")
        with patch.object(core_logic, "get_document_details", return_value=fake) as read:
            self.inbox.prepare(docx)
            self.assertTrue(read.call_args.kwargs["allow_cloud_ai"])
            self.inbox.prepare(pptx)
            self.assertFalse(read.call_args.kwargs["allow_cloud_ai"])

    def test_signature_change_during_extraction_requires_new_review(self):
        path = self.pdf()
        self.inbox.refresh()
        self.inbox.refresh()
        original = core_logic.get_basic_document_details
        def edited(source):
            details = original(source)
            source.write_bytes(source.read_bytes().replace(b"%%EOF", b"% changed while extracting\n%%EOF"))
            return details
        with patch.object(core_logic, "get_basic_document_details", side_effect=edited):
            with self.assertRaises(DownloadNotReady):
                self.inbox.prepare(path)
        self.assertTrue(path.exists())
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(path, "old-proposal.pdf", self.library)

    def test_signature_change_after_review_never_files_old_proposal(self):
        path = self.pdf()
        old = self.ready(path)
        path.write_bytes(path.read_bytes().replace(b"%%EOF", b"% changed after review\n%%EOF"))
        self.inbox.refresh()
        self.inbox.refresh()
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(path, "old-proposal.pdf", self.library)
        new = self.inbox.prepare(path)
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(path, "old-proposal.pdf", self.library, expected_signature=old["signature"])
        result = self.inbox.file_item(path, "new-proposal.pdf", self.library, expected_signature=new["signature"])
        self.assertTrue(Path(result["path"]).exists())

    def test_review_cache_avoids_repeat_online_or_extraction_calls(self):
        path = self.pdf()
        first = self.ready(path)
        with patch.object(core_logic, "get_basic_document_details", side_effect=AssertionError("use same review")):
            self.assertEqual(self.inbox.prepare(path), first)

    def test_explicit_retry_bypasses_review_cache(self):
        path = self.pdf()
        self.ready(path)
        original = core_logic.get_basic_document_details
        with patch.object(core_logic, "get_basic_document_details", wraps=original) as read:
            self.inbox.prepare(path)
            read.assert_not_called()
            self.inbox.prepare(path, force=True)
            read.assert_called_once()

    def test_filename_style_updates_cached_proposal_without_repeating_lookup(self):
        path = self.pdf()
        original = self.ready(path)
        self.inbox.save_config({"filename_format": "title_year_type"})
        with patch.object(core_logic, "get_basic_document_details", side_effect=AssertionError("format only")):
            updated = self.inbox.prepare(path)
        expected = core_logic.build_proposed_filename(original, ".pdf", filename_format="title_year_type")
        self.assertEqual(updated["proposed_filename"], expected)

    def test_naming_mode_change_requires_review_again_before_filing(self):
        path = self.pdf()
        self.ready(path)
        self.inbox.save_config({"naming_mode": "Automatic", "online_metadata_lookup_enabled": False})
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(path, "old-proposal.pdf", self.library)
        self.assertEqual(self.inbox.prepare(path)["source"], "Basic")

    def test_skip_persists_and_file_change_returns_it_to_review(self):
        path = self.pdf()
        self.ready(path)
        self.inbox.skip(path)
        reopened = InboxController(self.store, self.data)
        try:
            reopened.stable_seconds = 0
            reopened.refresh()
            self.assertEqual(reopened.list_items()[0]["status"], "skipped")
            path.write_bytes(path.read_bytes().replace(b"%%EOF", b"% changed\n%%EOF"))
            reopened.refresh()
            reopened.refresh()
            self.assertEqual(reopened.list_items()[0]["status"], "ready")
        finally:
            reopened.stop()

    def test_confirmed_filing_preserves_bytes_and_persists_metadata_job(self):
        path = self.pdf()
        expected = path.read_bytes()
        self.ready(path)
        result = self.inbox.file_item(path, "reviewed", self.library)
        self.assertFalse(path.exists())
        self.assertEqual(Path(result["path"]).read_bytes(), expected)
        self.assertEqual(result["digest"], hashlib.sha256(expected).hexdigest())
        self.assertTrue(result["is_pdf"])
        self.assertEqual(result["details"]["title"], "Synthetic Research Paper")
        self.assertEqual(self.inbox.list_items(), [])
        self.assertEqual(self.inbox.pending_filings(), [result])
        self.inbox.acknowledge_filing(result["job_id"])
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_word_powerpoint_and_legacy_ppt_file_with_original_extensions(self):
        word = self.incoming / "report.docx"
        document = Document()
        document.add_paragraph("Synthetic report")
        document.save(word)
        slides = self.incoming / "slides.pptx"
        _write_pptx(slides, ["Synthetic slides"])
        legacy = self.incoming / "legacy.ppt"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"Synthetic legacy binary")
        for path in (word, slides, legacy):
            with self.subTest(suffix=path.suffix):
                expected = path.read_bytes()
                self.ready(path)
                result = self.inbox.file_item(path, "filed" + path.suffix, self.library)
                self.assertFalse(result["is_pdf"])
                self.assertEqual(Path(result["path"]).suffix, path.suffix)
                self.assertEqual(Path(result["path"]).read_bytes(), expected)

    def test_filing_collision_retains_source_and_existing_file(self):
        path = self.pdf()
        existing = self.library / "already.pdf"
        existing.write_bytes(b"existing paper, do not overwrite")
        self.ready(path)
        with self.assertRaises(FileExistsError):
            self.inbox.file_item(path, existing.name, self.library)
        self.assertEqual(existing.read_bytes(), b"existing paper, do not overwrite")
        self.assertTrue(path.exists())
        self.assertEqual(self.inbox.pending_filings(), [])
        self.assertEqual(list(self.library.glob(".inbox-copy-*")), [])

    def test_concurrent_filings_never_overwrite_same_destination(self):
        first, second = self.pdf("first.pdf"), self.pdf("second.pdf")
        self.ready(first)
        self.ready(second)
        with ThreadPoolExecutor(max_workers=2) as workers:
            futures = [workers.submit(self.inbox.file_item, source, "same.pdf", self.library) for source in (first, second)]
            successes, errors = [], []
            for future in futures:
                try:
                    successes.append(future.result())
                except FileExistsError as error:
                    errors.append(error)
        self.assertEqual(len(successes), 1)
        self.assertEqual(len(errors), 1)
        self.assertEqual(sum(path.exists() for path in (first, second)), 1)
        self.assertEqual(len(self.inbox.pending_filings()), 1)

    def test_collision_created_during_copy_cannot_be_overwritten(self):
        path = self.pdf()
        self.ready(path)
        target = self.library / "late-collision.pdf"
        real_stat = shutil.copystat
        def another_writer(source, temporary):
            real_stat(source, temporary)
            target.write_bytes(b"another process created this paper")
        with patch("note_organizer.inbox.shutil.copystat", side_effect=another_writer):
            with self.assertRaises(FileExistsError):
                self.inbox.file_item(path, target.name, self.library)
        self.assertTrue(path.exists())
        self.assertEqual(target.read_bytes(), b"another process created this paper")
        self.assertEqual(list(self.library.glob(".inbox-copy-*")), [])

    def test_source_change_during_copy_cannot_be_confirmed(self):
        path = self.pdf()
        self.ready(path)
        copy = self.inbox._exclusive_copy
        def changed(source, target, auto_suffix=False):
            result = copy(source, target, auto_suffix)
            source.write_bytes(source.read_bytes().replace(b"%%EOF", b"% changed while copying\n%%EOF"))
            return result
        with patch.object(self.inbox, "_exclusive_copy", side_effect=changed):
            with self.assertRaises(DownloadNotReady):
                self.inbox.file_item(path, "cannot-file.pdf", self.library)
        self.assertTrue(path.exists())
        self.assertTrue((self.library / "cannot-file.pdf").exists())
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_changed_source_does_not_delete_an_edited_published_copy(self):
        source = self.pdf()
        self.ready(source)
        target = self.library / "edited-copy.pdf"
        copy = self.inbox._exclusive_copy
        def changed(path, destination, auto_suffix=False):
            result = copy(path, destination, auto_suffix)
            source.write_bytes(b"new download content")
            target.write_bytes(b"separate edits to the published copy")
            return result
        with patch.object(self.inbox, "_exclusive_copy", side_effect=changed):
            with self.assertRaisesRegex(DownloadNotReady, "Both copies were retained"):
                self.inbox.file_item(source, target.name, self.library)
        self.assertEqual(source.read_bytes(), b"new download content")
        self.assertEqual(target.read_bytes(), b"separate edits to the published copy")
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_add_changed_source_preserves_the_edited_inbox_copy(self):
        external = self.root / "external"
        external.mkdir()
        source = self.pdf(folder=external)
        target = self.incoming / source.name
        copy = self.inbox._exclusive_copy
        def changed(path, destination, auto_suffix=False):
            result = copy(path, destination, auto_suffix)
            source.write_bytes(b"new external content")
            target.write_bytes(b"edits to the inbox copy")
            return result
        with patch.object(self.inbox, "_exclusive_copy", side_effect=changed):
            with self.assertRaisesRegex(DownloadNotReady, "Both copies were retained"):
                self.inbox.add_files([source])
        self.assertEqual(source.read_bytes(), b"new external content")
        self.assertEqual(target.read_bytes(), b"edits to the inbox copy")

    def _journal_copy(self, source, target, job_id="guarded-recovery"):
        target.write_bytes(source.read_bytes())
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,digest,state,created_at) VALUES (?,?,?,?,?,'copied',?)",
                (job_id, canonical(source), canonical(target), file_signature(source),
                 hashlib.sha256(target.read_bytes()).hexdigest(), timestamp()))

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_final_filing_guard_blocks_concurrent_edits_and_replacements(self):
        for recovery in (False, True):
            with self.subTest(recovery=recovery):
                source = self.pdf(f"source-{recovery}.pdf")
                self.ready(source)
                expected = source.read_bytes()
                target = self.library / f"filed-{recovery}.pdf"
                replacement = self.root / f"replacement-{recovery}.pdf"
                replacement.write_bytes(b"another download")
                if recovery:
                    self._journal_copy(source, target)
                attempts = []
                def final_gap(kernel, handle):
                    operations = (
                        lambda: source.write_bytes(b"source editor"),
                        lambda: os.replace(replacement, source),
                        lambda: target.write_bytes(b"destination editor"),
                        lambda: os.replace(replacement, target),
                    )
                    with ThreadPoolExecutor(max_workers=1) as worker:
                        for operation in operations:
                            with self.assertRaises(PermissionError):
                                worker.submit(operation).result(timeout=5)
                            attempts.append(True)
                    remarkable_cloud._delete_guarded(kernel, handle)
                with patch("note_organizer.inbox._delete_guarded", side_effect=final_gap):
                    if recovery:
                        results = self.inbox.pending_filings()
                    else:
                        results = [self.inbox.file_item(source, target.name, self.library)]
                self.assertEqual(len(attempts), 4)
                self.assertEqual(len(results), 1)
                self.assertFalse(source.exists())
                self.assertEqual(target.read_bytes(), expected)
                self.assertEqual(replacement.read_bytes(), b"another download")
                self.inbox.acknowledge_filing(results[0]["job_id"])

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_open_source_writer_retains_both_copies_and_surfaces_recovery(self):
        for recovery in (False, True):
            with self.subTest(recovery=recovery):
                source = self.pdf(f"writer-source-{recovery}.pdf")
                self.ready(source)
                expected = source.read_bytes()
                target = self.library / f"writer-filed-{recovery}.pdf"
                if recovery:
                    self._journal_copy(source, target, f"source-writer-{recovery}")
                with source.open("r+b"):
                    if recovery:
                        self.assertEqual(self.inbox.pending_filings(), [])
                    else:
                        with self.assertRaisesRegex(DownloadNotReady, "open for editing"):
                            self.inbox.file_item(source, target.name, self.library)
                self.assertEqual(source.read_bytes(), expected)
                self.assertEqual(target.read_bytes(), expected)
                issues = [issue for issue in self.inbox.list_recovery_errors() if issue["source"] == canonical(source)]
                self.assertEqual(len(issues), 1)
                self.assertIn("retained for review", issues[0]["error"])
                if recovery:
                    self.assertIn(("inbox_recovery_error", issues[0]), self.events)

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_open_destination_writer_retains_both_complete_copies(self):
        for recovery in (False, True):
            with self.subTest(recovery=recovery):
                source = self.pdf(f"target-source-{recovery}.pdf")
                self.ready(source)
                expected = source.read_bytes()
                target = self.library / f"target-filed-{recovery}.pdf"
                editors = []
                copy = self.inbox._exclusive_copy
                def open_editor(path, destination, auto_suffix=False):
                    result = copy(path, destination, auto_suffix)
                    editors.append(target.open("r+b"))
                    return result
                if recovery:
                    self._journal_copy(source, target, f"target-writer-{recovery}")
                    editors.append(target.open("r+b"))
                try:
                    with patch.object(self.inbox, "_exclusive_copy", side_effect=open_editor):
                        if recovery:
                            self.assertEqual(self.inbox.pending_filings(), [])
                        else:
                            with self.assertRaisesRegex(DownloadNotReady, "open for editing"):
                                self.inbox.file_item(source, target.name, self.library)
                finally:
                    for editor in editors:
                        editor.close()
                self.assertEqual(source.read_bytes(), expected)
                self.assertEqual(target.read_bytes(), expected)
                self.assertEqual(self.inbox.pending_filings(), [])

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_source_hash_prevents_deletion_when_stat_signature_is_unchanged(self):
        for recovery in (False, True):
            with self.subTest(recovery=recovery):
                source = self.pdf(f"hash-source-{recovery}.pdf")
                self.ready(source)
                expected, signature = source.read_bytes(), file_signature(source)
                target = self.library / f"hash-checked-{recovery}.pdf"
                copy = self.inbox._exclusive_copy
                def changed(path, destination, auto_suffix=False):
                    result = copy(path, destination, auto_suffix)
                    source.write_bytes(b"changed source with colliding metadata")
                    return result
                if recovery:
                    self._journal_copy(source, target)
                    source.write_bytes(b"changed source with colliding metadata")
                actual_signature = file_signature
                def stale_signature(path):
                    return signature if Path(path) == source else actual_signature(path)
                with patch.object(self.inbox, "_exclusive_copy", side_effect=changed), \
                     patch("note_organizer.inbox.file_signature", side_effect=stale_signature):
                    if recovery:
                        self.assertEqual(self.inbox.pending_filings(), [])
                        self.assertIn("Inbox source changed", self.inbox.list_recovery_errors()[-1]["error"])
                    else:
                        with self.assertRaisesRegex(DownloadNotReady, "Inbox source changed"):
                            self.inbox.file_item(source, target.name, self.library)
                self.assertEqual(source.read_bytes(), b"changed source with colliding metadata")
                self.assertEqual(target.read_bytes(), expected)
                self.assertEqual(self.inbox.pending_filings(), [])

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_destination_change_before_guard_retains_the_complete_source(self):
        for recovery in (False, True):
            with self.subTest(recovery=recovery):
                source = self.pdf(f"destination-race-{recovery}.pdf")
                self.ready(source)
                expected = source.read_bytes()
                target = self.library / f"destination-race-{recovery}.pdf"
                if recovery:
                    self._journal_copy(source, target)
                guarded_file = remarkable_cloud._guarded_file
                def edit_before_guard(path, digest, action, item_label="PDF"):
                    if Path(path) == target:
                        target.write_bytes(b"destination changed before its handle opened")
                    return guarded_file(path, digest, action, item_label=item_label)
                with patch("note_organizer.inbox._guarded_file", side_effect=edit_before_guard):
                    if recovery:
                        self.assertEqual(self.inbox.pending_filings(), [])
                    else:
                        with self.assertRaisesRegex(DownloadNotReady, "filed document changed"):
                            self.inbox.file_item(source, target.name, self.library)
                self.assertEqual(source.read_bytes(), expected)
                self.assertEqual(target.read_bytes(), b"destination changed before its handle opened")
                self.assertEqual(self.inbox.pending_filings(), [])

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_recovery_delete_permission_failure_retries_without_losing_a_copy(self):
        source = self.pdf()
        self.ready(source)
        expected = source.read_bytes()
        target = self.library / "permission-recovery.pdf"
        self._journal_copy(source, target)
        self.events.clear()
        with patch("note_organizer.inbox._delete_guarded", side_effect=PermissionError("synthetic source deletion lock")):
            self.assertEqual(self.inbox.pending_filings(), [])
            issues = self.inbox.list_recovery_errors()
            self.assertEqual(len(issues), 1)
            self.assertIn("synthetic source deletion lock", issues[0]["error"])
            self.assertIn("retry automatically", issues[0]["error"])
            self.assertEqual(self.inbox.list_items()[0]["recovery_errors"], issues)
            self.assertEqual(self.inbox.pending_filings(), [])
            notifications = [payload for event, payload in self.events if event == "inbox_recovery_error"]
            self.assertEqual(notifications, issues)
        self.assertEqual(source.read_bytes(), expected)
        self.assertEqual(target.read_bytes(), expected)
        with self.store.connect() as db:
            row = db.execute("SELECT state,error FROM inbox_jobs WHERE id='guarded-recovery'").fetchone()
            self.assertEqual(row["state"], "copied")
            self.assertEqual(row["error"], issues[0]["error"])
        with patch("note_organizer.inbox._delete_guarded", side_effect=PermissionError("different temporary lock")):
            self.assertEqual(self.inbox.pending_filings(), [])
        notifications = [payload for event, payload in self.events if event == "inbox_recovery_error"]
        self.assertEqual(len(notifications), 2)
        self.assertIn("different temporary lock", notifications[-1]["error"])
        self.assertEqual(len(self.inbox.pending_filings()), 1)
        self.assertEqual(self.inbox.list_recovery_errors(), [])
        with self.store.connect() as db:
            row = db.execute("SELECT state,error FROM inbox_jobs WHERE id='guarded-recovery'").fetchone()
            self.assertEqual(row["state"], "confirmed")
            self.assertEqual(row["error"], "")
        self.assertFalse(source.exists())
        self.assertEqual(target.read_bytes(), expected)

    @unittest.skipUnless(os.name == "nt", "Windows mandatory sharing guards")
    def test_same_content_source_replacement_before_guard_requires_review(self):
        source = self.pdf()
        self.ready(source)
        expected = source.read_bytes()
        target = self.library / "same-content.pdf"
        replacement = self.root / "replacement.pdf"
        replacement.write_bytes(expected)
        guarded_file = remarkable_cloud._guarded_file
        def replace_before_guard(path, digest, action, item_label="PDF"):
            if Path(path) == source:
                os.replace(replacement, source)
            return guarded_file(path, digest, action, item_label=item_label)
        with patch("note_organizer.inbox._guarded_file", side_effect=replace_before_guard):
            with self.assertRaisesRegex(DownloadNotReady, "source changed"):
                self.inbox.file_item(source, target.name, self.library)
        self.assertEqual(source.read_bytes(), expected)
        self.assertEqual(target.read_bytes(), expected)
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_unsupported_platform_retains_copies_instead_of_unprotected_delete(self):
        source = self.pdf()
        self.ready(source)
        expected = source.read_bytes()
        target = self.library / "platform-copy.pdf"
        with patch("note_organizer.remarkable_cloud.os", SimpleNamespace(name="posix")):
            with self.assertRaisesRegex(DownloadNotReady, "requires Windows file sharing"):
                self.inbox.file_item(source, target.name, self.library)
        self.assertEqual(source.read_bytes(), expected)
        self.assertEqual(target.read_bytes(), expected)
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_new_destination_is_created_only_when_filing_is_confirmed(self):
        path = self.pdf()
        destination = self.library / "New field" / "Subtopic"
        self.ready(path)
        self.assertFalse(destination.exists())
        result = self.inbox.file_item(path, "paper.pdf", destination)
        self.assertTrue(destination.is_dir())
        self.assertEqual(Path(result["path"]).parent, destination)

    def test_destination_and_filename_validation_precedes_any_write(self):
        path = self.pdf()
        self.ready(path)
        for destination in (self.incoming, self.root / "Outside", self.data):
            with self.subTest(destination=destination), self.assertRaises(ValueError):
                self.inbox.file_item(path, "paper.pdf", destination)
        for name in ("../escape.pdf", "CON.pdf", "paper.docx", "bad:name.pdf"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.inbox.file_item(path, name, self.library)
        self.assertTrue(path.exists())
        self.assertEqual(list(self.library.iterdir()), [])

    def test_config_rejects_nested_and_managed_roots_without_side_effects(self):
        for config in ({"library_root": str(self.incoming / "Sorted")},
                       {"watch_folder": str(self.library / "ToSort")},
                       {"watch_folder": str(self.data / "downloads")},
                       {"library_root": str(self.data / "remarkable")}):
            with self.subTest(config=config), self.assertRaises(ValueError):
                self.inbox.save_config(config)
        self.assertEqual(self.inbox.config()["watch_folder"], str(self.incoming))
        self.assertFalse((self.incoming / "Sorted").exists())

    def test_add_files_copies_preserves_source_and_suffixes_name_collisions(self):
        external = self.root / "external"
        external.mkdir()
        source = self.pdf(folder=external)
        expected = source.read_bytes()
        existing = self.incoming / source.name
        existing.write_bytes(b"existing")
        targets = self.inbox.add_files([source])
        self.assertEqual(Path(targets[0]).name, "incoming (2).pdf")
        self.assertEqual(Path(targets[0]).read_bytes(), expected)
        self.assertEqual(source.read_bytes(), expected)
        self.assertEqual(existing.read_bytes(), b"existing")

    def test_add_files_rejects_managed_annotated_pdf_and_rmdoc(self):
        source = self.pdf(folder=self.library)
        self.store.save_remarkable_item("cloud-id", {"path": str(source), "state": "synced"})
        with self.assertRaises(ValueError):
            self.inbox.add_files([source])
        archive = self.library / "source.rmdoc"
        archive.write_bytes(b"archive")
        with self.assertRaises(ValueError):
            self.inbox.add_files([archive])
        self.assertEqual(list(self.incoming.iterdir()), [])

    def test_callback_failure_keeps_confirmed_filing_for_restart_handoff(self):
        path = self.pdf()
        self.ready(path)
        def offline(event, payload):
            if event == "inbox_filed":
                raise RuntimeError("UI was closing")
        self.inbox.post = offline
        with self.assertLogs(level="ERROR"):
            result = self.inbox.file_item(path, "filed.pdf", self.library)
        self.assertEqual(self.inbox.pending_filings(), [result])
        self.assertFalse(path.exists())

    def test_journal_recovery_only_removes_source_when_copy_is_verified(self):
        source = self.pdf()
        self.ready(source)
        target = self.library / "copied.pdf"
        target.write_bytes(source.read_bytes())
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,digest,state,created_at) VALUES (?,?,?,?,?,'copied',?)",
                ("test-job", canonical(source), canonical(target), file_signature(source), digest, timestamp()))
        self.inbox.refresh()
        self.assertFalse(source.exists())
        self.assertTrue(target.exists())
        self.assertEqual(self.inbox.pending_filings()[0]["job_id"], "test-job")

    def test_journal_recovery_preserves_changed_source_and_copy(self):
        source = self.pdf()
        signature = file_signature(source)
        target = self.library / "copied.pdf"
        target.write_bytes(source.read_bytes())
        source.write_bytes(source.read_bytes().replace(b"%%EOF", b"% new annotations\n%%EOF"))
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,digest,state,created_at) VALUES (?,?,?,?,?,'copied',?)",
                ("test-job", canonical(source), canonical(target), signature,
                 hashlib.sha256(target.read_bytes()).hexdigest(), timestamp()))
        self.inbox.refresh()
        self.assertTrue(source.exists())
        self.assertTrue(target.exists())
        self.assertEqual(self.inbox.pending_filings(), [])

    def test_startup_pending_filings_recovers_without_an_inbox_scan(self):
        source = self.pdf()
        prepared = self.ready(source)
        target = self.library / "recovered.pdf"
        target.write_bytes(source.read_bytes())
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,digest,details,state,created_at) VALUES (?,?,?,?,?,?,'copied',?)",
                ("startup-job", canonical(source), canonical(target), file_signature(source),
                 hashlib.sha256(target.read_bytes()).hexdigest(), json.dumps(prepared), timestamp()))
        fresh = InboxController(self.store, self.data)
        try:
            with patch.object(fresh, "refresh", side_effect=AssertionError("startup must not need a scan")):
                results = fresh.pending_filings()
            self.assertEqual([job["job_id"] for job in results], ["startup-job"])
            self.assertFalse(source.exists())
            self.assertTrue(target.exists())
            self.assertIsNone(fresh._thread)
            fresh.acknowledge_filing("startup-job")
            self.assertEqual(fresh.pending_filings(), [])
        finally:
            fresh.stop()

    def test_refresh_delivers_recovered_filing_only_once(self):
        source = self.pdf()
        self.ready(source)
        target = self.library / "recovered.pdf"
        target.write_bytes(source.read_bytes())
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,digest,state,created_at) VALUES (?,?,?,?,?,'copied',?)",
                ("refresh-job", canonical(source), canonical(target), file_signature(source),
                 hashlib.sha256(target.read_bytes()).hexdigest(), timestamp()))
        self.events.clear()
        self.inbox.refresh()
        self.inbox.refresh()
        delivered = [payload for event, payload in self.events if event == "inbox_filed"]
        self.assertEqual([job["job_id"] for job in delivered], ["refresh-job"])
        self.assertEqual(self.inbox.pending_filings(), delivered)

    def test_recovery_issues_remain_readable_with_and_without_the_inbox_source(self):
        source = self.pdf()
        self.ready(source)
        target = self.library / "ambiguous.pdf"
        target.write_bytes(source.read_bytes())
        with self.store.connect() as db:
            db.execute("INSERT INTO inbox_jobs (id,source,path,signature,state,created_at) VALUES (?,?,?,?,'copying',?)",
                ("ambiguous-job", canonical(source), canonical(target), file_signature(source), timestamp()))
        self.events.clear()
        self.inbox.refresh()
        errors = self.inbox.list_recovery_errors()
        self.assertEqual([row["job_id"] for row in errors], ["ambiguous-job"])
        self.assertIn("retained for review", errors[0]["error"])
        self.assertEqual(self.inbox.list_items()[0]["recovery_errors"], errors)
        self.assertTrue(source.exists())
        self.assertTrue(target.exists())
        self.assertEqual([payload for event, payload in self.events if event == "inbox_recovery_error"], errors)
        self.inbox.refresh()
        self.assertEqual(len([event for event, _ in self.events if event == "inbox_recovery_error"]), 1)
        source.unlink()
        self.inbox.refresh()
        self.assertEqual(self.inbox.list_items(), [])
        self.assertEqual(self.inbox.list_recovery_errors(), errors)
        self.assertTrue(target.exists())

    def test_key_removal_and_review_invalidation_rollback_together(self):
        source = self.pdf()
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"), \
             patch("note_organizer.inbox._unprotect_for_current_windows_user", return_value="synthetic-own-key"):
            self.inbox.save_config({"api_key": "synthetic-own-key", "gemini_enabled": True})
        self.ready(source)
        saved, cached = self.store.preference("inbox"), self.inbox.config()
        generation = self.inbox._metadata_generation
        with self.store.connect() as db:
            db.execute("""CREATE TRIGGER reject_review_change BEFORE UPDATE ON inbox_items
                          BEGIN SELECT RAISE(ABORT, 'synthetic invalidation failure'); END""")
        self.events.clear()
        with self.assertRaises(sqlite3.IntegrityError):
            self.inbox.save_config({"api_key": "", "gemini_enabled": False})
        self.assertEqual(self.store.preference("inbox"), saved)
        self.assertEqual(self.inbox.config(), cached)
        self.assertEqual(self.inbox._metadata_generation, generation)
        self.assertEqual(self.inbox.list_items()[0]["status"], "review")
        self.assertEqual(self.events, [])
        with self.store.connect() as db:
            db.execute("DROP TRIGGER reject_review_change")
        self.inbox.save_config({"api_key": "", "gemini_enabled": False})
        self.assertFalse(self.inbox.config()["api_key_present"])
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")

    def test_other_preferences_failure_rolls_back_key_and_metadata(self):
        source = self.pdf()
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"):
            self.inbox.save_config({"api_key": "synthetic-own-key"})
        self.ready(source)
        before = self.store.preference("inbox")
        def fail_other_preferences(db):
            db.execute("INSERT INTO preferences(key,value) VALUES('pdf_viewer', ?)", (json.dumps({"viewer": "system"}),))
            raise sqlite3.OperationalError("synthetic preferences failure")
        with self.assertRaises(sqlite3.OperationalError):
            self.inbox.save_config({"api_key": ""}, persist_preferences=fail_other_preferences)
        self.assertEqual(self.store.preference("inbox"), before)
        self.assertIsNone(self.store.preference("pdf_viewer"))
        self.assertTrue(self.inbox.config()["api_key_present"])
        self.assertEqual(self.inbox.list_items()[0]["status"], "review")

    def test_watcher_failure_after_commit_is_reported_as_saved(self):
        with patch.object(self.inbox, "start", side_effect=RuntimeError("synthetic watcher failure")):
            result = self.inbox.save_config({"filename_format": "title_year_type"})
        self.assertEqual(result["filename_format"], "title_year_type")
        self.assertEqual(self.store.preference("inbox")["filename_format"], "title_year_type")
        errors = [payload["error"] for event, payload in self.events if event == "inbox_error"]
        self.assertIn("Preferences were saved", errors[-1])

    def test_naming_change_during_filename_build_cannot_publish_stale_review(self):
        source = self.pdf()
        self.inbox.refresh()
        self.inbox.refresh()
        entered, release = threading.Event(), threading.Event()
        original = core_logic.build_proposed_filename
        def delayed(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise TimeoutError("synthetic preparation barrier")
            return original(*args, **kwargs)
        try:
            with patch.object(core_logic, "build_proposed_filename", side_effect=delayed), ThreadPoolExecutor(max_workers=1) as workers:
                result = workers.submit(self.inbox.prepare, source)
                self.assertTrue(entered.wait(2))
                self.inbox.save_config({"naming_mode": "Automatic", "online_metadata_lookup_enabled": False})
                release.set()
                with self.assertRaisesRegex(DownloadNotReady, "Naming settings changed"):
                    result.result(timeout=3)
        finally:
            release.set()
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")
        self.assertEqual(self.inbox.list_items()[0]["details"], {})
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(source, "stale.pdf", self.library)
        reviewed = self.inbox.prepare(source)
        filed = self.inbox.file_item(source, reviewed["proposed_filename"], self.library, expected_signature=reviewed["signature"])
        self.assertTrue(Path(filed["path"]).exists())

    def test_sorter_settings_preview_is_read_only_and_never_exposes_key(self):
        path = self.root / "old-settings.json"
        original = json.dumps({"watch_folder": "OldInbox", "sorted_folder": "OldTopics", "api_key": "synthetic-secret",
                              "watch_and_launch_enabled": True, "naming_mode": "AI", "filename_format": "author_year_title"})
        path.write_text(original)
        with patch("note_organizer.inbox._unprotect_for_current_windows_user", side_effect=AssertionError("preview must not decrypt")):
            result = self.inbox.preview_sorter_settings(path)
        self.assertEqual(result["watch_folder"], str(self.root / "OldInbox"))
        self.assertEqual(result["library_root"], str(self.root / "OldTopics"))
        self.assertEqual(result["naming_mode"], "Automatic")
        self.assertFalse(result["enabled"])
        self.assertTrue(result["api_key_present"])
        self.assertNotIn("synthetic-secret", json.dumps(result))
        self.assertEqual(path.read_text(), original)
        self.assertEqual(self.inbox.config()["watch_folder"], str(self.incoming))

    def test_explicit_import_keeps_watching_off_and_key_is_opt_in(self):
        path = self.root / "old-settings.json"
        data = {"watch_folder": str(self.incoming), "sorted_folder": str(self.library),
                "enabled": True, "watch_and_launch_enabled": True, "gemini_enabled": True,
                "api_key": "synthetic-secret"}
        path.write_text(json.dumps(data))
        result = self.inbox.import_sorter_settings(path)
        self.assertFalse(result["enabled"])
        self.assertFalse(result["api_key_present"])
        self.assertFalse(result["gemini_enabled"])
        self.assertIsNone(self.inbox._thread)
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="protected-test-key"):
            result = self.inbox.import_sorter_settings(path, include_key=True)
        self.assertTrue(result["api_key_present"])
        self.assertFalse(result["gemini_enabled"])
        self.assertNotIn("synthetic-secret", json.dumps(result))
        self.assertNotIn("api_key", self.store.preference("inbox"))
        self.assertEqual(json.loads(path.read_text()), data)

    def test_secret_setting_is_protected_and_clear_is_explicit(self):
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"):
            result = self.inbox.save_config({"api_key": "synthetic-secret"})
        self.assertTrue(result["api_key_present"])
        self.assertNotIn("synthetic-secret", json.dumps(self.store.preference("inbox")))
        self.inbox.save_config({"naming_mode": "Basic"})
        self.assertTrue(self.inbox.config()["api_key_present"])
        self.inbox.save_config({"api_key": ""})
        self.assertFalse(self.inbox.config()["api_key_present"])

    def test_remove_saved_key_persists_without_changing_folders_flags_or_watcher(self):
        with patch("note_organizer.inbox._protect_for_current_windows_user", return_value="encrypted"), \
             patch.object(self.inbox, "refresh", return_value=[]):
            self.inbox.save_config({"api_key": "synthetic-secret", "enabled": True,
                                   "allow_cloud_ai_for_word_documents": True})
            original = self.inbox.config()
            watcher = self.inbox._thread
            result = self.inbox.remove_api_key()
            self.assertFalse(result["api_key_present"])
            self.assertEqual({key: value for key, value in result.items() if key != "api_key_present"},
                             {key: value for key, value in original.items() if key != "api_key_present"})
            self.assertIs(self.inbox._thread, watcher)
            self.assertTrue(watcher.is_alive())
        fresh = InboxController(self.store, self.data)
        try:
            fresh.initialize()
            self.assertFalse(fresh.config()["api_key_present"])
            self.assertEqual(fresh.config()["watch_folder"], str(self.incoming))
            self.assertTrue(fresh.config()["enabled"])
            self.assertTrue(fresh.config()["allow_cloud_ai_for_word_documents"])
            self.assertNotIn("api_key_protected", self.store.preference("inbox"))
        finally:
            fresh.stop()

    def test_remove_saved_key_clears_legacy_plaintext_and_invalidates_smart_review(self):
        path = self.pdf()
        self.inbox.save_config({"naming_mode": "Automatic", "online_metadata_lookup_enabled": False})
        self.ready(path)
        saved = self.store.preference("inbox")
        saved["api_key"] = "synthetic-legacy-key"
        self.store.set_preference("inbox", saved)
        result = self.inbox.remove_api_key()
        self.assertFalse(result["api_key_present"])
        self.assertNotIn("api_key", self.store.preference("inbox"))
        self.assertEqual(self.inbox.list_items()[0]["status"], "ready")
        with self.assertRaises(DownloadNotReady):
            self.inbox.file_item(path, "old-proposal.pdf", self.library)

    def test_preparation_async_posts_results_from_worker_and_stop_is_repeatable(self):
        path = self.pdf()
        self.inbox.refresh()
        self.inbox.refresh()
        threads = []
        self.inbox.post = lambda event, _: threads.append((event, threading.current_thread().name))
        result = self.inbox.prepare_async(path).result(timeout=5)
        self.assertEqual(result["source"], "Basic")
        self.assertTrue(any(event == "inbox_prepared" and name.startswith("Inbox-review") for event, name in threads))
        self.inbox.stop(wait=False)
        self.inbox.stop()
        with self.assertRaises(RuntimeError):
            self.inbox.prepare_async(path)

    def test_enabled_watching_owns_one_thread_and_stops_cleanly(self):
        with patch.object(self.inbox, "refresh", return_value=[]):
            self.inbox.save_config({"enabled": True})
            first = self.inbox._thread
            self.assertTrue(self.inbox.start())
            self.assertIs(self.inbox._thread, first)
            self.inbox.save_config({"enabled": False})
            self.assertFalse(first.is_alive())
            self.assertFalse(self.inbox.start())


if __name__ == "__main__":
    unittest.main()
