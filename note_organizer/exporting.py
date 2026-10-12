"""Atomic, reproducible Word exports with preservation of external edits."""
import base64
from datetime import datetime
from html import unescape
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import tempfile
from itertools import groupby

import pymupdf as fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .storage import digest_file, entry_sort_key
from .pages import page_description
from .extraction import snapshot
from .remarkable_cloud import CloudError, publish_staged


# Notebook layout changes can refresh exports without re-extracting PDF annotations.
EXPORT_FORMAT_REVISION = 4


def file_uri(path, page):
    return Path(path).resolve().as_uri() + f"#page={page}"


def _source_link(entry, page=None):
    """Notebook links must outlive rebuildable reading previews.

    An explicitly empty path means there is no verified durable counterpart.
    Older entries without the field retain their original source link.
    """
    path = entry.get("source_link_path", entry["path"])
    target_page = entry.get("source_link_page", entry["page"]) if page is None else page
    if not path or not isinstance(target_page, int) or target_page < 1:
        return None
    return file_uri(path, target_page)


def _source_hyperlink(paragraph, label, entry, page=None):
    target = _source_link(entry, page)
    if target:
        hyperlink(paragraph, label, target)
    else:
        paragraph.add_run(label + " · source page needs review")


def _link_guidance(document, entries):
    if any("source_link_path" in entry for entry in entries):
        document.add_paragraph("For verified papers, PDF links open the saved original paper. "
                               "Older tablet copies keep links to their retained reading PDFs until a successful refresh. "
                               "Open PDF in Notable shows the "
                               "combined PC and tablet notes; Export annotated PDF saves that view as a separate PDF.",
                               style="Caption")


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245A81")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _add_page_context_link(document, entry):
    paragraph = document.add_paragraph(style="Caption")
    _source_hyperlink(paragraph, f"Full page context — PDF page {entry['page']}", entry)


def _preview_rect(entry):
    """Only saved crop coordinates establish the screenshot's actual region."""
    try:
        rect = fitz.Rect(entry["crop"])
        if not rect.is_empty and not rect.is_infinite and all(abs(v) < 1e6 for v in rect):
            return rect
    except (KeyError, TypeError, ValueError):
        pass
    return None


def _shared_preview_region(a, b):
    if a.contains(b) or b.contains(a):
        return True
    # Ignore slivers caused by the crop's outer padding. Compare original
    # windows, never an expanding union that could absorb unrelated notes.
    inner_a = fitz.Rect(a.x0 + 6, a.y0 + 6, a.x1 - 6, a.y1 - 6)
    inner_b = fitz.Rect(b.x0 + 6, b.y0 + 6, b.x1 - 6, b.y1 - 6)
    overlap = inner_a & inner_b
    return (not overlap.is_empty and
            overlap.get_area() >= .20 * min(inner_a.get_area(), inner_b.get_area()))


def consolidate_drawing_previews(entries):
    """Share notebook images without merging or changing stored annotations.

    Reuse a covering saved preview, or render unions from an unchanged source.
    If no faithful union can be made, preserve the original saved previews.
    """
    entries = [dict(entry) for entry in entries]
    pages = {}
    for index, entry in enumerate(entries):
        if entry["kind"] == "drawing" and entry.get("image"):
            rect = _preview_rect(entry)
            if rect is not None:
                pages.setdefault((entry["path"], entry["page"], entry.get("source_digest")), []).append((index, rect))
    pending = {}

    def share(members, image):
        # Place the shared context after all of its drawing notes. Every note's
        # quote, comment, transcription and detailed metadata is still emitted.
        for index in members:
            entries[index]["image"] = ""
        entries[max(members)]["image"] = image

    for (path, page, _), windows in pages.items():
        parents = list(range(len(windows)))
        def root(i):
            while parents[i] != i:
                parents[i] = parents[parents[i]]
                i = parents[i]
            return i
        for i, (_, rect) in enumerate(windows):
            for j in range(i):
                if _shared_preview_region(rect, windows[j][1]):
                    parents[root(i)] = root(j)
        groups = {}
        for i, window in enumerate(windows):
            groups.setdefault(root(i), []).append(window)
        for group in groups.values():
            if len(group) < 2:
                continue
            members = [index for index, _ in group]
            union = fitz.Rect(group[0][1])
            for _, rect in group[1:]:
                union |= rect
            covering = next((index for index, rect in group if rect.contains(union)), None)
            if covering is not None:
                share(members, entries[covering]["image"])
            else:
                pending.setdefault(path, []).append((page, union, members))

    for path, plans in pending.items():
        # Stored crops may outlive a source file or an edit that has not yet
        # been indexed. Verify the source before and after rendering; never
        # substitute a screenshot of a different PDF revision.
        digests = {entries[index].get("source_digest") for _, _, members in plans for index in members}
        if len(digests) != 1 or not next(iter(digests)):
            continue
        expected = next(iter(digests))
        try:
            if digest_file(path) != expected:
                continue
            with fitz.open(path) as pdf:
                images = [(members, snapshot(pdf[page - 1], rect)) for page, rect, members in plans]
            if digest_file(path) != expected:
                continue
        except (OSError, RuntimeError, ValueError, IndexError):
            continue
        for members, image in images:
            share(members, image)
    return entries


def _build_detailed_notebook(title, entries):
    doc = Document()
    doc.core_properties.version = str(EXPORT_FORMAT_REVISION)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("Annotations collected from your PDFs, with source links and visual context. "
                      "This notebook is regenerated when the source annotations change. "
                      "Use an editable copy for your own writing.")
    doc.add_paragraph(f"Generated {datetime.now():%d %B %Y at %H:%M}. {len(entries)} annotations.")
    _link_guidance(doc, entries)
    if not entries:
        doc.add_paragraph("No annotations have been extracted yet.")
    previous = None
    paper_groups = {key: list(paper) for key, paper in groupby(entries, key=_paper_group_key)}
    for e in entries:
        group_key = _paper_group_key(e)
        if group_key != previous:
            previous = group_key
            doc.add_heading(_entry_paper_title(e), 1)
            for source in _paper_sources(paper_groups[group_key]):
                path = source.get("source_link_path") or source["path"]
                p = doc.add_paragraph()
                _source_hyperlink(p, Path(path).name, source, 1)
                doc.add_paragraph(str(Path(path).parent), style="Caption")
                meta = source.get("metadata", {})
                if meta.get("author"):
                    doc.add_paragraph(f"Author: {meta['author']}")
                if meta.get("doi"):
                    p = doc.add_paragraph("DOI: ")
                    hyperlink(p, meta["doi"], "https://doi.org/" + meta["doi"])
                if source.get("missing"):
                    doc.add_paragraph("Source unavailable. Previously extracted notes are retained.")
                for warning in meta.get("warnings", []):
                    doc.add_paragraph(warning)
        p = doc.add_paragraph(style="Heading 2")
        _source_hyperlink(p, page_description(e), e)
        p.add_run(f"  |  {e['kind'].capitalize()}" + (f"  |  {e['color']}" if e.get("color") else ""))
        created = re.match(r"(?:D:)?(\d{4})(\d{2})(\d{2})", e.get("created", ""))
        if created:
            doc.add_paragraph("Annotation date: " + "-".join(created.groups()), style="Caption")
        if e.get("quote"):
            p = doc.add_paragraph()
            p.add_run("Highlighted text: ").bold = True
            p.add_run(e["quote"])
        if e.get("comment"):
            p = doc.add_paragraph()
            p.add_run("Comment: ").bold = True
            p.add_run(e["comment"])
        if e.get("image") and e["kind"] == "drawing":
            raw = base64.b64decode(e["image"])
            with Image.open(BytesIO(raw)) as img:
                width, height = img.size
            # Keep page crops readable without forcing tall images beyond a page.
            inches = min(6.7, max(1.5, width / 144), 7.2 * width / max(height, 1))
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(BytesIO(raw), width=Inches(inches))
            _add_page_context_link(doc, e)
        if e.get("ocr"):
            p = doc.add_paragraph()
            p.add_run("Suggested transcription (verify against image): ").bold = True
            p.add_run(e["ocr"])
    return doc


def clean_text(value):
    """PDF strings can contain control characters that Word's XML rejects."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", str(value or ""))


def ordered_entries(entries):
    groups = {}
    for entry in sorted(entries, key=entry_sort_key):
        groups.setdefault(_paper_group_key(entry), []).append(entry)
    result = []
    for group in groups.values():
        # Keep each variant's pages together. Unlinked exports retain their
        # established path/page/reading-order sort exactly.
        if group[0].get("logical_paper_id"):
            group.sort(key=lambda entry: (
                str(entry.get("source_link_path") or entry["path"]).casefold(),
                "" if entry.get("source_link_path") else str(entry.get("document_id", "")),
                "" if entry.get("source_link_path") else str(entry.get("source_digest", "")),
                entry.get("source_link_page") if isinstance(entry.get("source_link_page"), int) else entry["page"],
                *entry_sort_key(entry)[2:]))
        result.extend(group)
    return result


def _paper_group_key(entry):
    paper_id = entry.get("logical_paper_id")
    return ("paper", paper_id) if paper_id else ("path", entry["path"])


def _paper_sources(entries):
    sources = {}
    for entry in entries:
        sources.setdefault(entry.get("source_link_path") or entry["path"], entry)
    return list(sources.values())


def _source_page_key(entry):
    if entry.get("source_link_path") and isinstance(entry.get("source_link_page", entry["page"]), int):
        return (entry["source_link_path"], entry.get("source_link_page", entry["page"]))
    return (entry["path"], entry.get("document_id"), entry.get("source_digest"), entry["page"])


def _entry_paper_title(entry):
    metadata = entry.get("metadata", {})
    if entry.get("logical_paper_title"):
        metadata = {**metadata, "title": entry["logical_paper_title"]}
    return paper_title(entry["path"], metadata)


def paper_title(path, metadata):
    title = unescape(metadata.get("title", "")).strip()
    # Publishers sometimes set the title to a typesetting job ID or page range.
    if not title or re.search(r"\.indd$|\d+\.\.\d+$|^[A-Z]{1,6}[_-]\d{4,}", title):
        return Path(path).stem.replace("_", " ")
    return title


def build_notebook(title, entries, layout="Reading"):
    # Preserve the stored source verbatim; sanitize only the XML representation.
    entries = [{**e, **{key: clean_text(e.get(key, "")) for key in
                ("quote", "comment", "ocr", "label", "created", "color")},
                **{key: clean_text(e[key]) for key in ("pdf_label", "journal_page") if key in e},
                **{key: clean_text(e[key]) for key in ("logical_paper_title",) if key in e},
                "metadata": {k: clean_text(v) if isinstance(v, str) else v
                             for k, v in e.get("metadata", {}).items()}}
               for e in ordered_entries(entries)]
    entries = consolidate_drawing_previews(entries)
    title = clean_text(title)
    if layout == "Detailed":
        return _build_detailed_notebook(title, entries)
    doc = Document()
    doc.core_properties.version = str(EXPORT_FORMAT_REVISION)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    styles = doc.styles
    for name in ("Normal", "Quote"):
        styles[name].font.name = "Georgia"
        styles[name].font.size = Pt(11)
        styles[name].font.color.rgb = RGBColor.from_string("20252A")
        styles[name].paragraph_format.line_spacing = 1.15
        styles[name].paragraph_format.space_after = Pt(8)
        styles[name].paragraph_format.widow_control = True
    styles["Quote"].paragraph_format.left_indent = Inches(.18)
    styles["Quote"].paragraph_format.right_indent = Inches(.1)
    styles["Quote"].font.italic = False
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 11)):
        styles[name].font.name = "Segoe UI"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.space_before = Pt(16 if name == "Heading 1" else 9)
        styles[name].paragraph_format.space_after = Pt(7)
    styles["Caption"].font.name = "Segoe UI"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.color.rgb = RGBColor.from_string("59636D")
    doc.add_paragraph(title, "Title")
    if not entries:
        doc.add_paragraph("No notes yet.")
        return doc
    doc.add_paragraph("Selected passages and reading notes", "Subtitle")
    _link_guidance(doc, entries)
    for _, paper_iter in groupby(entries, key=_paper_group_key):
        paper = list(paper_iter)
        doc.add_heading(_entry_paper_title(paper[0]), 1)
        sources = _paper_sources(paper)
        for source_entry in sources:
            path = source_entry.get("source_link_path") or source_entry["path"]
            meta = source_entry.get("metadata", {})
            source = doc.add_paragraph(style="Caption")
            if meta.get("author"):
                source.add_run(meta["author"] + "  ·  ")
            _source_hyperlink(source, Path(path).name, source_entry, 1)
            if meta.get("doi"):
                source.add_run("  ·  ")
                hyperlink(source, "DOI", "https://doi.org/" + meta["doi"])
            if source_entry.get("missing"):
                doc.add_paragraph("Source unavailable. Previously saved notes are retained.", "Caption")
        page_key = _source_page_key if paper[0].get("logical_paper_id") else lambda entry: entry["page"]
        for _, page_iter in groupby(paper, key=page_key):
            page_entries = list(page_iter)
            source_entry = page_entries[0]
            heading = doc.add_paragraph(style="Heading 2")
            _source_hyperlink(heading, page_description(source_entry), source_entry)
            if len(sources) > 1:
                heading.add_run("  ·  " + Path(source_entry["path"]).name)
            for e in page_entries:
                if e["quote"]:
                    quote = doc.add_paragraph(e["quote"], "Quote")
                    # Keep the source passage with its attached comment where possible.
                    quote.paragraph_format.keep_with_next = bool(e["comment"])
                elif e["kind"] in {"highlight", "underline", "strikeout", "squiggly"}:
                    doc.add_paragraph("Highlighted text unavailable; open the linked source PDF.", "Caption")
                if e["comment"]:
                    p = doc.add_paragraph()
                    p.add_run("Note: ").bold = True
                    p.add_run(e["comment"])
                image = e.get("image")
                needs_image = e["kind"] == "drawing"
                if image and needs_image:
                    raw = base64.b64decode(image)
                    with Image.open(BytesIO(raw)) as img:
                        width, height = img.size
                    inches = min(6.8, max(1.5, width/144), 7.3*width/max(height, 1))
                    preview = doc.add_paragraph()
                    preview.paragraph_format.keep_with_next = True
                    preview.add_run().add_picture(BytesIO(raw), width=Inches(inches))
                    _add_page_context_link(doc, e)
                if e["ocr"]:
                    doc.add_paragraph("Transcription to check: " + e["ocr"], "Caption")
    return doc


def export_notebook(path, title, entries, expected_hash="", layout="Reading"):
    """Publish a verified staged notebook while preserving external edits."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    original_hash = digest_file(path) if path.exists() else None
    fd, temp = tempfile.mkstemp(prefix=".pdf-notes-", suffix=".docx", dir=path.parent)
    os.close(fd)
    backup = None
    try:
        build_notebook(title, entries, layout=layout).save(temp)
        # Verify the zip is readable before replacing a working notebook.
        Document(temp)
        generated_hash = digest_file(temp)
        if original_hash is not None and original_hash != expected_hash:
            backup_dir = path.parent / "Notebook backups"
            backup_dir.mkdir(exist_ok=True)
            # A locked Word file may be retried repeatedly: preserve one copy per revision.
            backup = backup_dir / f"{path.stem}.{original_hash}.docx"
            if not backup.exists():
                shutil.copy2(path, backup)
            if digest_file(backup) != original_hash:
                raise OSError("Notebook changed while preserving edits; retry when Word is closed.")
        now_hash = digest_file(path) if path.exists() else None
        if now_hash != original_hash:
            raise OSError("Notebook changed during export; retry when Word is closed.")
        try:
            retained = publish_staged(temp, path, expected_digest=original_hash,
                                      require_absent=original_hash is None, item_label="notebook")
        except (CloudError, OSError) as failure:
            raise OSError(f"Notebook export paused for {path}: {failure} Close Word and retry.") from failure
        # Record the generated revision even if an editor changes the published
        # target immediately afterward; that edit must receive a backup next time.
        return generated_hash, str(retained or backup) if retained or backup else ""
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
