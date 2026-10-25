"""Validate the graphical plugin in a disposable Zotero profile and library.

The helper extension below is test instrumentation, never shipped to users.
It authorizes only the synthetic profile and reports the actual rendered UI.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import uuid

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tools.build_zotero_support import build
from note_organizer.zotero_pilot import LocalZotero, PilotError
from note_organizer.zotero_worker import Worker
from tests.test_cloud_outbox import MemoryCloud, pdf_bytes


BOOTSTRAP = """
async function startup({rootURI}) {
  const config = JSON.parse(await (await fetch(rootURI + 'config.json')).text());
  Zotero.uiReadyPromise.then(async () => {
    if (Zotero.DataDirectory.dir !== config.data) throw new Error('Unexpected test library');
    Zotero.Server.LocalAPI._promptForAuthorization = async name => ({allow: name === 'Notable Zotero Pilot', remember: true});
    await IOUtils.writeUTF8(config.ready, JSON.stringify({version: Zotero.version, server: Zotero.Server.LocalAPI.getServerID(), pid: Services.appinfo.processID}));
    while (!(await IOUtils.exists(config.done))) {
      const path = PathUtils.join(PathUtils.parent(config.ready), 'command.json');
      if (await IOUtils.exists(path)) {
        const command = JSON.parse(await IOUtils.readUTF8(path)); await IOUtils.remove(path);
        let response = {id: command.id};
        try {
          const ui = Zotero.Notable;
          if (!ui) throw new Error('Notable UI did not initialize');
          if (command.action === 'open') {
            ui.open(); const deadline = Date.now() + 120000;
            while ((!ui.process || !Object.keys(ui.latest?.types || {}).length || !ui.view?.document.getElementById('app')?.querySelector('h1')) && Date.now() < deadline) {
              if (ui.view?.document.querySelector('[role="alert"]')) break;
              await Zotero.Promise.delay(100);
            }
          } else if (command.action === 'connect') await ui.action('connect');
          else if (command.action === 'tab') { ui.tab = command.tab; ui.render(); }
          else if (command.action === 'folders') await ui.loadFolders();
          else if (command.action === 'refresh') { await ui.refresh(); ui.render(); }
          else if (command.action === 'render') { ui.latest = command.snapshot; ui.tab = command.tab || 'review'; ui.render(); }
          else if (command.action === 'collection_names') {
            const collection = await Zotero.Collections.getByLibraryAndKeyAsync(Zotero.Libraries.userLibraryID, command.key);
            if (!collection || collection.deleted) throw Error('Missing synthetic collection');
            const originalName = collection.name;
            ui.latest = command.snapshot; ui.folderCatalog = command.folders;
            ui.tab = 'review'; ui.render();
            const doc = ui.view.document;
            const filename = [...doc.querySelectorAll('label')].find(label => label.firstElementChild?.textContent === 'Filename').querySelector('input');
            filename.value = 'My unsaved name.pdf'; filename.dispatchEvent(new ui.view.Event('input'));
            const destination = doc.querySelector('select[data-folder-destination]');
            const connection = ui.filingConnection(), rpc = ui.rpc;
            const ensure = (value, message) => {if (!value) throw Error(message);};
            ui.rpc = () => {throw Error('Renaming unexpectedly contacted the worker');};
            try {
              collection.name = 'Literature'; await collection.saveTx();
              ensure(destination.selectedOptions[0].textContent === 'Zotero: Literature → reMarkable: Literature', 'Open destination kept old name');
              ensure(filename.isConnected && filename.value === 'My unsaved name.pdf', 'Rename replaced the active review form');
              ensure(ui.filingConnection() === connection && destination.value === command.key, 'Rename changed folder identity');
              ui.tab = 'setup'; ui.render();
              ensure(doc.getElementById('notable-connected-zotero').textContent === 'Zotero: Literature', 'Setup kept old collection name');
              response.passed = ['native rename updates open destination', 'review edit survives rename', 'connection and destination keys unchanged', 'setup shows renamed collection'];
              ui.rpc = rpc;
              ui.view.close(); ui.view = null; ui.open();
              const deadline = Date.now() + 15000;
              while (!ui.view?.document.getElementById('app')?.querySelector('h1') && Date.now() < deadline) await Zotero.Promise.delay(50);
              await ui.refresh(); ui.tab = 'setup'; ui.render();
              ensure(ui.latest.config.zotero_path === 'Literature', 'Reopened worker status restored stale name');
              ensure(ui.view.document.getElementById('notable-connected-zotero').textContent === 'Zotero: Literature', 'Reopened Setup kept old name');
              response.passed.push('reopened workspace refreshes stale saved name');
            } finally {
              ui.rpc = rpc; collection.name = originalName; await collection.saveTx();
            }
          }
          else if (command.action === 'layout') {
            ui.busy = !!command.operation; ui.operation = command.operation || null;
            ui.changingFolder = !!command.editFolder; ui.setupDirty = false; ui.setupDraft = null;
            ui.folderPreview = command.folderPreview || null; ui.disconnectPreview = command.disconnectPreview || null;
            if (command.snapshot) ui.latest = command.snapshot;
            if (command.folders) ui.folderCatalog = command.folders;
            if (command.tagEditor) ui.tagEditor = command.tagEditor;
            ui.tab = command.tab || 'setup'; ui.render();
            if (command.expandDetails) for (const detail of ui.view.document.querySelectorAll('details')) detail.open = true;
            ui.view.resizeTo(command.width || 1040, command.height || 760);
            ui.view.document.documentElement.style.colorScheme = command.theme || 'light';
            await Zotero.Promise.delay(300);
            const doc = ui.view.document, rect = node => {
              const r = node.getBoundingClientRect();
              return {x: r.x, y: r.y, width: r.width, height: r.height, right: r.right, bottom: r.bottom};
            };
            response.layout = {
              width: ui.view.innerWidth, scrollWidth: doc.documentElement.scrollWidth,
              root: doc.documentElement.localName,
              buttons: [...doc.querySelectorAll('button')].filter(b => b.getClientRects().length).map(b => {
                const icon = b.querySelector('.button-icon');
                return {text: b.textContent, ...rect(b), icon: !!icon, iconWidth: icon?.getBoundingClientRect().width || 0,
                  iconMask: icon ? ui.view.getComputedStyle(icon).maskImage : 'none'};
              }),
              checkboxDisplays: [...doc.querySelectorAll('.checkbox-field')].map(n => ui.view.getComputedStyle(n).display),
              pairingWidth: doc.querySelector('.pairing-code')?.getBoundingClientRect().width,
              selects: [...doc.querySelectorAll('select')].map(s => ({...rect(s), appearance: ui.view.getComputedStyle(s).appearance, multiple: s.multiple, size: s.size, options: s.options.length,
                optionRects: [...s.options].slice(0, 4).map(o => ({text: o.textContent, ...rect(o), display: ui.view.getComputedStyle(o).display}))}))
            };
          }
          else if (command.action === 'picker') {
            const win = ui.view, doc = win.document;
            const select = doc.querySelectorAll('select')[command.index || 0];
            select.scrollIntoView({block: 'center'}); select.focus();
            const r = select.getBoundingClientRect();
            win.windowUtils.sendMouseEvent('mousedown', r.right - 16, r.y + r.height / 2, 0, 1, 0);
            win.windowUtils.sendMouseEvent('mouseup', r.right - 16, r.y + r.height / 2, 0, 1, 0);
            await Zotero.Promise.delay(300);
            response.popups = [...doc.querySelectorAll('menupopup')].map(p => ({state: p.state, background: win.getComputedStyle(p).backgroundColor,
              appearance: win.getComputedStyle(p).appearance, items: p.children.length, width: p.getBoundingClientRect().width, height: p.getBoundingClientRect().height}));
            response.valueBefore = select.value;
            if (command.choose !== undefined) {
              const popup = [...doc.querySelectorAll('menupopup')].find(p => p.state === 'open');
              const item = popup?.querySelectorAll('menuitem')[command.choose];
              if (!item) throw new Error('No open native option menu');
              item.click();
              await Zotero.Promise.delay(100);
            }
            if (command.close !== false) for (const p of doc.querySelectorAll('menupopup')) p.hidePopup();
            response.valueAfter = select.value;
          }
          else if (command.action === 'upgrade') {
            if (command.packaged && Services.prefs.prefHasUserValue('notable.developmentWorker')) Services.prefs.clearUserPref('notable.developmentWorker');
            const {AddonManager} = ChromeUtils.importESModule('resource://gre/modules/AddonManager.sys.mjs');
            const file = Cc['@mozilla.org/file/local;1'].createInstance(Ci.nsIFile); file.initWithPath(command.path);
            const install = await AddonManager.getInstallForFile(file);
            await install.install();
            const addon = await AddonManager.getAddonByID('notable-companion@local.invalid');
            response.addon = {id: addon.id, name: addon.name, version: addon.version};
          }
          else if (command.action === 'folder_ui') {
            const original = ui.rpc, saved = {latest: ui.latest, folders: ui.folderCatalog};
            const oldRoot = '11111111-1111-4111-8111-111111111111', newRoot = '22222222-2222-4222-8222-222222222222';
            let state = {config: {zotero_root: 'CNCT2345', tablet_root: oldRoot, default_folder: oldRoot,
              zotero_path: 'Tablet papers', tablet_path: 'Previous folder', automatic: true},
              status: {}, types: {}, reviews: [], links: [], mappings: {}, collections: [{key: 'CNCT2345', name: 'Tablet papers', parentCollection: false}],
              connection: {linked_papers: 3, previous: []}};
            response.requests = [];
            ui.rpc = async (method, params = {}) => {
              response.requests.push({method, params});
              if (method === 'pause_sync') { state.config.sync_paused = true; return JSON.parse(JSON.stringify(state)); }
              if (method === 'status') return JSON.parse(JSON.stringify(state));
              if (method === 'preview_connection') return {token: 'preview', proposed: {...params, zotero_path: 'Tablet papers', tablet_path: 'Literature'},
                linked_papers: 3, pending_reviews: 1, restored_papers: 0, zotero_folders: [], tablet_folders: [], excluded_zotero_folders: ['Old topic'], excluded_tablet_folders: []};
              if (method === 'configure') { state.config = {...params, zotero_path: 'Tablet papers', tablet_path: 'Literature', automatic: false, sync_paused: true}; return JSON.parse(JSON.stringify(state)); }
              if (method === 'preview_disconnect') return {token: 'disconnect', linked_papers: 3, pending_reviews: 1};
              if (method === 'disconnect_folder') { state.config.tablet_root = ''; state.config.automatic = false; return JSON.parse(JSON.stringify(state)); }
              throw new Error('Unexpected UI request: ' + method);
            };
            const until = Date.now() + 15000;
            const waitFor = async condition => { while (!condition() && Date.now() < until) await Zotero.Promise.delay(30); if (!condition()) throw new Error('Folder UI did not settle'); };
            const click = label => { const button = [...ui.view.document.querySelectorAll('button')].find(b => b.textContent === label); if (!button || button.disabled) throw new Error('Unavailable button: ' + label); button.click(); };
            try {
              ui.clearFolderEditor(); ui.latest = JSON.parse(JSON.stringify(state)); ui.tab = 'setup';
              ui.folderCatalog = {collections: state.collections, folders: [{id: oldRoot, name: 'Previous folder', parent: ''}, {id: newRoot, name: 'Literature', parent: ''}]}; ui.render();
              click('Change folder'); await waitFor(() => !ui.pausing && ui.view.document.querySelectorAll('main select').length === 3);
              const select = ui.view.document.querySelectorAll('main select')[1]; select.value = newRoot;
              select.dispatchEvent(new ui.view.Event('change', {bubbles: true})); select.dispatchEvent(new ui.view.Event('input', {bubbles: true}));
              if (!ui.view.document.getElementById('notable-sync').disabled) throw new Error('Sync enabled during unsaved change');
              click('Review folder change'); await waitFor(() => !!ui.folderPreview && !ui.busy);
              response.previewShown = ui.view.document.querySelector('main').textContent.includes('Old Zotero folders kept out');
              if (response.requests.some(r => ['configure', 'sync'].includes(r.method))) throw new Error('Preview changed a connection');
              click('Save folder change'); await waitFor(() => !ui.changingFolder && !ui.busy);
              response.savedWithoutSync = !response.requests.some(r => r.method === 'sync') && ui.latest.config.tablet_root === newRoot;
              click('Disconnect folder'); await waitFor(() => !!ui.disconnectPreview && !ui.busy && !ui.pausing);
              click('Cancel'); await waitFor(() => !ui.disconnectPreview);
              if (!ui.latest.config.tablet_root) throw new Error('Cancel disconnected the folder');
              click('Disconnect folder'); await waitFor(() => !!ui.disconnectPreview && !ui.busy && !ui.pausing);
              click('Confirm disconnect'); await waitFor(() => !ui.disconnectPreview && !ui.busy);
              response.disconnected = !ui.latest.config.tablet_root && ui.view.document.getElementById('notable-sync').disabled;
              const entry = Zotero.getMainWindow().document.getElementById('notable-open');
              response.toolbar = {label: entry.getAttribute('label'), tooltip: entry.getAttribute('tooltiptext'), image: entry.style.listStyleImage};
              if (!response.toolbar.image.includes('tablet.svg') || !response.toolbar.tooltip.startsWith('Open reMarkable')) throw new Error('Ambiguous toolbar entry');
            } finally { ui.rpc = original; ui.clearFolderEditor(); ui.latest = saved.latest; ui.folderCatalog = saved.folders; ui.tab = 'setup'; ui.render(); }
          }
          else if (command.action === 'filing_ui') {
            const saved = {latest: ui.latest, tab: ui.tab, rpc: ui.rpc, queue: ui.filingQueue, paused: ui.filingQueuePaused};
            const requests = []; let release, job;
            const gate = new Promise(resolve => { release = resolve; });
            let state = JSON.parse(JSON.stringify(command.snapshot));
            const ensure = (value, reason) => { if (!value) throw Error(reason); };
            const waitFor = async check => { const until = Date.now() + 15000; while (!check() && Date.now() < until) await Zotero.Promise.delay(30); ensure(check(), 'Filing UI did not settle'); };
            const field = label => [...ui.view.document.querySelectorAll('label.field')].find(node => node.firstElementChild.textContent === label)?.querySelector('input, select');
            const edit = (label, value) => { const input = field(label); ensure(input && !input.disabled, 'Field locked: ' + label); input.value = value; input.dispatchEvent(new ui.view.Event('input')); input.dispatchEvent(new ui.view.Event('change')); };
            try {
              ui.filingQueue = []; ui.filingQueuePaused = false; ui.queueError = ''; ui.reviewDrafts.clear(); ui.selected = null;
              ui.latest = state; ui.tab = 'review'; ui.render();
              await waitFor(() => [...(field('Main paper in Zotero (required)')?.options || [])].some(option => option.value === command.mainKey));
              ensure(field('Document type').value === 'supporting_information', 'Supporting information was not selected');
              edit('Main paper in Zotero (required)', command.mainKey);
              edit('Title', 'Reviewed supplement title'); edit('Filename', 'Supplement_SI.pdf');
              ui.rpc = async (method, params) => {
                requests.push({method, params});
                if (method === 'approve_routine') await gate;
                if (method === 'approve') state.reviews = state.reviews.filter(row => row.id !== params.identity);
                return JSON.parse(JSON.stringify(state));
              };
              job = ui.action('approve_routine', {});
              const approve = [...ui.view.document.querySelectorAll('button')].find(button => button.textContent === 'Add to filing queue');
              ensure(approve && !approve.disabled, 'Cannot queue during a batch'); approve.click();
              await waitFor(() => ui.filingQueue.length === 1); await ui.queueSave;
              const queued = JSON.parse(await IOUtils.readUTF8(PathUtils.join(ui.statePath, 'filing-queue.json'))).jobs[0];
              ensure(queued.params.main_reference === command.mainKey && queued.params.filename === 'Supplement_SI.pdf', 'Queued choices were not saved');
              ensure(queued.params.main_reference_metadata.title, 'Main citation snapshot missing');
              ensure(requests.length === 1, 'Approval ran concurrently with batch');
              edit('Title', 'Keep this next-paper draft');
              release(); await job; job = null;
              await waitFor(() => !ui.drainingFilings && !ui.busy && !ui.filingQueue.length);
              ensure(requests.map(row => row.method).join(',') === 'approve_routine,approve', 'Filing must not start a whole-library sync');
              ensure(field('Title').value === 'Keep this next-paper draft', 'Background completion lost the next draft');
              response.passed = ['SI choice and native main-paper search', 'editable review during filing', 'durable queued choices', 'one writer at a time', 'main citation snapshot', 'draft retained through filing', 'filing does not start whole-library sync'];
            } finally {
              release(); if (job) await job;
              ui.rpc = saved.rpc; ui.latest = saved.latest; ui.tab = saved.tab;
              ui.filingQueue = saved.queue; ui.filingQueuePaused = saved.paused; ui.queueError = ''; await ui.saveFilingQueue(); ui.render();
            }
          }
          else if (command.action === 'activity_ui') {
            const saved = {latest: ui.latest, tab: ui.tab, rpc: ui.rpc, launch: Zotero.launchFile,
              process: ui.process, selected: ui.selected, all: ui.showAllReviews};
            const requests = [], launches = [], wires = [];
            let release;
            const gate = new Promise(resolve => { release = resolve; });
            let job;
            const ensure = (condition, text) => { if (!condition) throw new Error(text); };
            const click = text => {
              const button = [...ui.view.document.querySelectorAll('button')].find(b => b.textContent === text);
              ensure(button && !button.disabled, 'Unavailable during work: ' + text); button.click();
            };
            const settle = () => Zotero.Promise.delay(100);
            try {
              ui.latest = command.snapshot; ui.tab = 'import'; ui.showAllReviews = false; ui.render();
              ui.schedule(); ensure(!ui.timer, 'Automatic check scheduled while workspace open');
              ui.rpc = async (method, params) => { requests.push(method); if (method === 'scan_library') return gate; if (method === 'status') return command.snapshot; throw new Error('Unexpected request: ' + method); };
              await ui.action('sync', {automatic: true});
              ensure(!requests.length, 'Automatic check ran while workspace open');
              Zotero.launchFile = file => { launches.push(file.path); };
              job = ui.action('scan_library', {root: command.source}); await settle();
              ui.message('Scanning file 12 of 537: Example paper.pdf'); ui.operation.started -= 65000; ui.updateActivity();
              const doc = ui.view.document;
              ensure(!doc.getElementById('activity').hidden, 'No working indicator');
              ensure(doc.getElementById('activity-progress').value === 12 && doc.getElementById('activity-progress').max === 537, 'Progress count missing');
              ensure(doc.getElementById('activity-time').textContent.includes('1 min'), 'Elapsed time missing');
              ensure([...doc.querySelectorAll('nav button')].every(b => !b.disabled), 'Navigation locked during scan');
              click('Review'); await settle();
              ensure(ui.busy && ui.tab === 'review', 'Browsing interrupted the job');
              ensure(![...doc.querySelectorAll('button')].find(b => b.textContent === 'Add to filing queue').disabled, 'Cannot queue a review during scan');
              ensure([...doc.querySelectorAll('.review-form input, .review-form select')].every(field => !field.disabled), 'Review fields locked during work');
              click('Open paper'); await settle();
              ensure(launches.length === 1 && launches[0] === command.source, 'Open paper did not use the saved local file');
              ensure(requests.length === 1, 'Opening a paper queued behind the worker or changed a library');
              click('Browse all papers'); await settle();
              ensure(doc.querySelectorAll('aside button').length === 2, 'Routine papers cannot be browsed');
              release(command.snapshot); await job; job = null;
              ensure(doc.getElementById('activity').hidden && !ui.busy, 'Working indicator stuck after completion');
              ui.operation = {automatic: true, started: Date.now(), updated: Date.now(), detail: 'Checking…', title: 'Automatic sync'};
              ui.busy = true; ui.process = {stdin: {write: async value => { wires.push(JSON.parse(value)); }}};
              await ui.yieldBackground();
              ensure(wires.length === 1 && wires[0].method === 'cancel', 'Opening workspace cannot yield background work');
              ensure(ui.latest.config.automatic && !ui.latest.config.sync_paused, 'Yield disabled the saved sync preference');
              ui.busy = false; ui.operation = null;
              const view = ui.view; ui.view = null;
              ensure(ui.automaticAllowed(), 'Automatic sync cannot resume after workspace closes');
              ui.schedule(); ensure(!!ui.timer, 'No closed-workspace timer');
              ui.view = view; ui.schedule(); ensure(!ui.timer, 'Opening workspace failed to suppress timer');
              response.passed = ['automatic checks wait while workspace open', 'navigation stays available during work',
                'visible progress and elapsed time', 'local paper opens without approval or import',
                'routine papers can be browsed', 'working indicator clears on completion',
                'background work yields without disabling automatic preference', 'automatic checks resume after workspace closes'];
            } finally {
              release(command.snapshot); if (job) await job;
              ui.rpc = saved.rpc; Zotero.launchFile = saved.launch; ui.process = saved.process;
              ui.busy = false; ui.operation = null; ui.latest = saved.latest; ui.tab = saved.tab;
              ui.selected = saved.selected; ui.showAllReviews = saved.all; ui.render(); ui.schedule();
            }
          }
          else if (command.action === 'review_list') {
            const saved = {latest: ui.latest, tab: ui.tab, rpc: ui.rpc};
            const ensure = (value, reason) => { if (!value) throw new Error(reason); };
            const settle = () => Zotero.Promise.delay(200);
            try {
              ui.latest = command.snapshot; ui.tab = 'review'; ui.selected = null; ui.showAllReviews = true; ui.reviewScrollTop = 0;
              ui.render(); ui.view.resizeTo(1040, 900); await settle();
              const doc = ui.view.document;
              const sizes = () => { const list = doc.querySelector('.review-layout > aside'), detail = list.nextElementSibling;
                return {list: list.getBoundingClientRect().height, detail: detail.getBoundingClientRect().height,
                  scroll: list.scrollHeight, visible: list.clientHeight}; };
              response.before = sizes();
              ensure(Math.abs(response.before.list - response.before.detail) < 2, 'List does not match detail height');
              ensure(response.before.scroll > response.before.visible * 2, 'Long list does not scroll independently');
              const pageTop = ui.view.scrollY;
              let button = doc.querySelector('aside button');
              button.dispatchEvent(new ui.view.KeyboardEvent('keydown', {key: 'End', bubbles: true})); await settle();
              ensure(ui.selected === command.snapshot.reviews.at(-1).id, 'End key did not select the last paper');
              let list = doc.querySelector('aside');
              ensure(list.scrollTop > 0 && Math.abs(ui.view.scrollY - pageTop) < 2, 'Selecting a paper scrolled the whole page');
              const selected = list.querySelector('[aria-pressed="true"]'), selectedBox = selected.getBoundingClientRect(), listBox = list.getBoundingClientRect();
              ensure(selectedBox.top >= listBox.top - 1 && selectedBox.bottom <= listBox.bottom + 1, 'Selected paper not visible');
              const position = list.scrollTop;
              ui.render(); await settle();
              ensure(Math.abs(doc.querySelector('aside').scrollTop - position) < 2, 'List position reset during render');
              doc.querySelector('.paper-tag-proposal').open = true; await settle();
              response.expanded = sizes();
              ensure(Math.abs(response.expanded.list - response.expanded.detail) < 2, 'List failed to follow expanded details');
              const requests = [];
              ui.rpc = async (method, params) => {
                requests.push({method, params});
                if (method === 'suggest_paper_tags') return {proposal: command.proposal};
                if (method === 'save_paper_tags') return {saved: true, deferred_until_filing: true};
                if (method === 'status') return command.snapshot;
                throw new Error('Unexpected tag UI request');
              };
              [...doc.querySelectorAll('button')].find(b => b.textContent === 'Find keywords and suggest tags').click(); await settle();
              const checks = [...doc.querySelector('.paper-tag-proposal').querySelectorAll('input[type="checkbox"]')];
              ensure(checks.length && checks.every(input => !input.checked), 'Suggestions applied without selection');
              checks[0].checked = true;
              [...doc.querySelectorAll('button')].find(b => b.textContent === 'Use these tags when filing').click(); await settle();
              ensure(requests.find(row => row.method === 'save_paper_tags')?.params.selected.length === 1, 'Tag selection not respected');
              response.tags = sizes(); ensure(Math.abs(response.tags.list - response.tags.detail) < 2, 'Tag controls broke matching heights');
              ui.view.resizeTo(640, 900); await settle();
              const narrow = sizes();
              ensure(narrow.list <= ui.view.innerHeight * .36 + 2 && narrow.scroll > narrow.visible, 'Narrow list is unbounded');
              response.narrow = narrow;
              const icon = await (await fetch('chrome://notable/content/icons/tablet.svg')).text();
              ensure(icon.includes('M9 16V8') && !icon.includes('<line'), 'Toolbar icon still contains a dot');
              response.passed = ['review list matches details height', 'long list scrolls independently', 'keyboard selection stays visible',
                'list position persists', 'expanded details and tag controls resize list', 'suggested tags require a selection',
                'narrow list remains bounded', 'toolbar uses rectangle and R without dot'];
            } finally { ui.rpc = saved.rpc; ui.latest = saved.latest; ui.tab = saved.tab; ui.selected = null; ui.showAllReviews = false; ui.view.resizeTo(1040, 760); ui.render(); }
          }
          else if (command.action === 'reindex') await ui.action('reindex', {key: command.key});
          else if (command.action === 'rpc') { response.result = await ui.action(command.method, command.params || {}); }
          else if (command.action === 'reader_tags') {
            const item = await Zotero.Items.getByLibraryAndKeyAsync(Zotero.Libraries.userLibraryID, command.context.key);
            const reader = await Zotero.Reader.open(item.id, {pageIndex: command.context.page});
            const until = Date.now() + 30000;
            while ((!reader._internalReader?._state?.primaryViewStats?.pagesCount
                || !reader._iframeWindow.document.querySelector('.notable-page-tag')) && Date.now() < until) await Zotero.Promise.delay(100);
            const button = reader._iframeWindow.document.querySelector('.notable-page-tag');
            if (!button) throw new Error('Reader tag control missing');
            const original = ui.rpc; response.requests = [];
            ui.rpc = async function(method, params) {
              if (method === 'tag_context') { response.requests.push({method, params}); return command.context; }
              if (method === 'set_page_tags') {
                response.requests.push({method, params});
                return {...command.context, page_tags: params.tags, paper_tags: params.paper_tags};
              }
              return original.call(this, method, params);
            };
            try {
              button.click();
              while (!ui.tagEditor && Date.now() < until) await Zotero.Promise.delay(100);
              if (!ui.tagEditor) throw new Error('Reader page editor did not open');
              response.heading = ui.view.document.querySelector('main h2')?.textContent;
              const inputs = [...ui.view.document.querySelectorAll('main input')];
              inputs[0].value = 'Native paper tag'; inputs[1].value = 'Native page tag';
              [...ui.view.document.querySelectorAll('button')].find(b => b.textContent === 'Save tags').click();
              while (!response.requests.some(r => r.method === 'set_page_tags') && Date.now() < until) await Zotero.Promise.delay(100);
              while (ui.busy && Date.now() < until) await Zotero.Promise.delay(100);
              response.readerPage = reader._internalReader._state.primaryViewStats.pageIndex;
              response.readerButton = button.textContent;
            } finally { ui.rpc = original; }
          }
          else if (command.action === 'snapshot') {
            const win = command.mainWindow ? Zotero.getMainWindow() : ui.view, canvas = win.document.createElementNS('http://www.w3.org/1999/xhtml', 'canvas');
            canvas.width = win.innerWidth; canvas.height = win.innerHeight;
            if (command.scrollY !== undefined) {
              win.scrollTo(0, command.scrollY);
              win.document.querySelector('body').scrollTop = command.scrollY;
            }
            await Zotero.Promise.delay(100);
            const ctx = canvas.getContext('2d');
            if (ctx.drawWindow) ctx.drawWindow(win, 0, win.scrollY, canvas.width, canvas.height, 'white');
            else {
              const image = await win.browsingContext.currentWindowGlobal.drawSnapshot({x: 0, y: 0, width: canvas.width, height: canvas.height}, 1, 'white');
              ctx.drawImage(image, 0, 0); image.close();
            }
            const raw = atob(canvas.toDataURL('image/png').split(',')[1]);
            await IOUtils.write(PathUtils.join(PathUtils.parent(config.ready), command.filename || 'workspace.png'), Uint8Array.from(raw, c => c.charCodeAt(0)));
            response.width = canvas.width; response.height = canvas.height;
          }
          else if (command.action === 'replay') {
            const record = await Zotero.Items.getByLibraryAndKeyAsync(Zotero.Libraries.userLibraryID, command.key);
            await record.loadAllData(); response.annotations = record.isPDFAttachment() ? record.getAnnotations().length : 0;
          } else throw new Error('Unsupported test command');
          response.message = ui.view?.document.getElementById('message')?.textContent;
          response.buttons = Array.from(ui.view?.document.querySelectorAll('button') || [], b => b.textContent);
          response.title = ui.view?.document.title;
          response.worker = !!ui.process;
        } catch (error) { response.error = error.message; response.stack = error.stack; }
        await IOUtils.writeUTF8(PathUtils.join(PathUtils.parent(config.ready), 'response.json'), JSON.stringify(response));
      }
      await Zotero.Promise.delay(150);
    }
    Services.startup.quit(Ci.nsIAppStartup.eAttemptQuit);
  });
}
function install() {} function shutdown() {} function uninstall() {}
"""


def command(root, action, **values):
    identity = str(uuid.uuid4())
    (root / 'command.tmp').write_text(json.dumps({'id': identity, 'action': action, **values}), encoding='utf-8')
    (root / 'command.tmp').replace(root / 'command.json')
    deadline = time.monotonic() + 135
    while time.monotonic() < deadline:
        path = root / 'response.json'
        if path.exists():
            try:
                result = json.loads(path.read_text(encoding='utf-8'))
            except json.JSONDecodeError:
                # The test extension may still be writing this response.
                # Keep the same command identity and existing overall timeout.
                pass
            else:
                if result.get('id') == identity:
                    if result.get('error'):
                        raise RuntimeError(result['error'] + '\n' + result.get('stack', ''))
                    return result
        time.sleep(.15)
    raise RuntimeError('Plugin UI command timed out')


def connection_checks(client, root, checked):
    """Native Zotero library + unchanged MemoryCloud; no production accounts."""
    client.authorize()
    import copy
    from unittest.mock import patch
    cloud = MemoryCloud()
    old = cloud.create_folder('Previous folder', '', str(uuid.uuid4()))
    new = cloud.create_folder('Literature', '', str(uuid.uuid4()))
    child = cloud.create_folder('Old topic', old.id, str(uuid.uuid4()), old.id)
    client.create('collection', {'key': 'CNCT2345', 'name': 'Connection test papers'})
    client.create('item', {'key': 'CNCTPR23', 'itemType': 'journalArticle', 'title': 'Connection preservation', 'collections': ['CNCT2345']})
    client.create('item', {'key': 'CNCTSR23', 'itemType': 'attachment', 'parentItem': 'CNCTPR23', 'title': 'Main PDF', 'contentType': 'application/pdf',
                           'linkMode': 'imported_file', 'filename': 'Connection.pdf'})
    path = root / 'Connection original.pdf'; path.write_bytes(pdf_bytes())
    client.upload('CNCTSR23', path, '')
    worker = Worker(root / 'connection-worker', zotero_factory=lambda: LocalZotero(int(client.endpoint.split(':')[2].split('/')[0]),
                    grant_store=client.grant_store, interactive=False), cloud_factory=lambda: cloud)
    worker.config = {'server_id': client.server_id, 'zotero_root': 'CNCT2345', 'tablet_root': old.id, 'default_folder': old.id, 'ocr': False, 'automatic': True}
    worker.save()
    with patch.object(worker.review, 'observe_ready', return_value=True): worker.sync()
    row = next(r for r in worker.review.rows() if r['identity'] == 'CNCTSR23')
    worker.approve(row['id'], row['digest'], 'Connection paper.pdf', 'Connection preservation', 'journal_article', 'CNCT2345')
    worker.sync()
    documents = worker.state.read('documents'); identity = next(iter(documents))
    documents[identity]['page_tags'] = {'saved-page-uuid': ['Keep page tag']}; worker.state.write('documents', documents)
    original_state = worker.state.directory
    before_catalog = client.review_action('catalog')
    attachment_bytes = {r['data']['key']: Path(r['path']).read_bytes() for r in before_catalog['attachments']}
    cloud_before = {r.id: r.hash for r in cloud.library()}
    token = worker.preview_connection(zotero_root='CNCT2345', tablet_root=new.id, default_folder=new.id)['token']
    worker.configure(zotero_root='CNCT2345', tablet_root=new.id, default_folder=new.id, confirmation=token)
    assert worker.config['sync_paused'] and not worker.config['automatic']
    assert client.review_action('catalog') == before_catalog
    assert {r.id: r.hash for r in cloud.library()} == cloud_before
    checked('native_folder_change_saves_without_sync_or_changes_to_either_library')
    worker.sync()
    assert not [r for r in cloud.library() if r.parent == new.id]
    assert worker.state.read('documents', {}) == {}
    assert all(client.file_path(key).read_bytes() == data for key, data in attachment_bytes.items())
    checked('native_new_scope_keeps_old_mirrored_folders_and_linked_papers_out_without_deleting_them')
    token = worker.preview_connection(zotero_root='CNCT2345', tablet_root=old.id, default_folder=old.id)['token']
    worker.configure(zotero_root='CNCT2345', tablet_root=old.id, default_folder=old.id, confirmation=token)
    assert worker.state.directory == original_state and worker.state.read('documents') == documents
    checked('native_return_to_old_folders_restores_original_links_page_tags_and_recovery_paths')
    token = worker.preview_disconnect()['token']; worker.disconnect_folder(token)
    assert not worker.config['tablet_root'] and not worker.status()['links']
    assert client.review_action('catalog') == before_catalog
    assert {r.id: r.hash for r in cloud.library()} == cloud_before
    checked('native_disconnect_keeps_papers_notes_collections_and_tablet_bytes_unchanged')
    # The packaged helper itself must also retain and disconnect saved state.
    command(root, 'rpc', method='pause_sync')
    preview = command(root, 'rpc', method='preview_disconnect')['result']
    disconnected = command(root, 'rpc', method='disconnect_folder', params={'confirmation': preview['token']})['result']
    assert not disconnected['config']['tablet_root'] and disconnected['config']['sync_paused']
    checked('packaged_helper_pauses_and_disconnects_without_cloud_access')
    ui = command(root, 'folder_ui')
    assert ui['previewShown'] and ui['savedWithoutSync'] and ui['disconnected'], ui
    checked('native_buttons_change_review_save_cancel_and_disconnect_without_an_implicit_sync')
    checked('native_toolbar_entry_uses_a_tablet_icon_and_an_open_workspace_tooltip')


def native_checks(client, root, checked):
    """Real Zotero API, the actual cloud writer against in-memory blob CAS."""
    cloud = MemoryCloud()
    croot = cloud.create_folder('Tablet literature', '', str(uuid.uuid4()))
    client.authorize()
    client.create('collection', {'key': 'SYNC2345', 'name': 'reMarkable'})
    client.create('collection', {'key': 'TPIC2345', 'name': 'Research topic'})
    client.create('item', {'key': 'PARENT23', 'itemType': 'thesis', 'title': 'Synthetic dissertation', 'collections': ['TPIC2345'], 'university': 'Example University'})
    source = root / 'Incoming original.pdf'; source.write_bytes(pdf_bytes())
    client.create('item', {'key': 'SURCE234', 'itemType': 'attachment', 'parentItem': 'PARENT23', 'title': 'Main text', 'contentType': 'application/pdf', 'linkMode': 'imported_file', 'filename': 'Original.pdf'})
    client.upload('SURCE234', source, '')
    catalog = client.review_action('catalog')
    assert catalog['parents']['PARENT23']['university'] == 'Example University'
    assert Path(catalog['attachments'][0]['path']).exists()
    checked('native_catalog_uses_actual_local_paths_and_preserves_thesis_metadata')
    # The same encrypted permission grant can be used by the private worker.
    worker = Worker(root / 'review-worker', zotero_factory=lambda: LocalZotero(int(client.endpoint.split(':')[2].split('/')[0]), grant_store=client.grant_store, interactive=False), cloud_factory=lambda: cloud)
    worker.config = {'server_id': client.server_id, 'zotero_root': 'SYNC2345', 'tablet_root': croot.id, 'default_folder': croot.id, 'ocr': False}
    worker.sync(); time.sleep(2.1); worker.sync()
    row = worker.review.rows()[0]
    assert not worker.state.read('documents', {})
    checked('connector_or_manual_import_waits_for_naming_and_destination_review')
    worker.approve(row['id'], row['digest'], 'Reviewed dissertation.pdf', 'Synthetic dissertation', 'thesis_dissertation', 'SYNC2345', ['TPIC2345'])
    assert client.file_path('SURCE234').name == 'Reviewed dissertation.pdf'
    worker.sync()
    links = worker.state.read('documents'); assert len(links) == 1
    cloud_id, link = next(iter(links.items()))
    assert set(client.get('item', 'PARENT23')['collections']) == {'SYNC2345', 'TPIC2345'}
    assert len(client.children('PARENT23')) == 1
    checked('reviewed_native_rename_and_single_tablet_home_with_multiple_collections')
    worker.sync(); assert len(worker.state.read('documents')) == 1
    checked('repeated_sync_does_not_duplicate_documents_or_reviews')
    client.create('collection', {'key': 'MVCL2345', 'name': 'Moved papers', 'parentCollection': 'SYNC2345'})
    worker.sync()
    mapped_folder = worker.state.read('folders')['MVCL2345']['cloud']
    current = next(r for r in cloud.library() if r.id == cloud_id)
    cloud.update_item(cloud_id, expected_hash=current.hash, name='Tablet renamed', parent=mapped_folder, scope_id=croot.id)
    worker.sync()
    assert client.file_path('SURCE234').name == 'Tablet renamed.pdf'
    assert set(client.get('item', 'PARENT23')['collections']) == {'MVCL2345', 'TPIC2345'}
    client.review_action('rename', {'key': 'SURCE234', 'filename': 'PC renamed.pdf', 'md5': client.file_md5('SURCE234')})
    parent = client.get('item', 'PARENT23'); parent['collections'] = ['SYNC2345', 'TPIC2345']; client.update(parent)
    worker.sync()
    current = next(r for r in cloud.library() if r.id == cloud_id)
    assert current.name == 'PC renamed' and current.parent == croot.id
    checked('PDF_rename_and_move_follow_both_directions_without_losing_research_collections')
    current = next(r for r in cloud.library() if r.id == cloud_id)
    cloud.update_item(cloud_id, expected_hash=current.hash, parent='trash', scope_id=croot.id)
    worker.sync(); decision = worker.review.deletions()[0]
    assert not client.get('item', 'PARENT23').get('deleted')
    assert decision['payload']['restore_parent'] == croot.id
    worker.delete_decision(decision['id'], 'restore')
    worker.sync(); assert not worker.review.deletions()
    checked('confirmed_tablet_trash_waits_and_explicit_restore_resumes_same_link')
    # The whole reference has real notes and SI. PDF-only must preserve them.
    client.create('item', {'key': 'NTES2345', 'itemType': 'note', 'parentItem': 'PARENT23', 'note': '<p>Retain this research note</p>'})
    client.create('item', {'key': 'SUPP2345', 'itemType': 'attachment', 'parentItem': 'PARENT23', 'title': 'SI', 'contentType': 'application/pdf', 'linkMode': 'imported_file', 'filename': 'SI.pdf'})
    client.upload('SUPP2345', source, '')
    current = next(r for r in cloud.library() if r.id == cloud_id)
    cloud.update_item(cloud_id, expected_hash=current.hash, parent='trash', scope_id=croot.id)
    worker.sync(); decision = worker.review.deletions()[0]
    # Mutating a note after the displayed inventory invalidates that decision.
    note = client.get('item', 'NTES2345'); note['note'] += '<p>New note</p>'; client.update(note)
    try:
        worker.delete_decision(decision['id'], 'pdf_only')
        raise AssertionError('Stale deletion was accepted')
    except PilotError:
        pass
    assert not client.get('item', 'SURCE234').get('deleted')
    checked('changed_notes_block_stale_deletion_without_trashing_a_PDF')
    # A refused transaction restores the pending decision automatically.
    assert worker.review.deletions()[0]['status'] == 'pending'
    worker.sync(); decision = worker.review.deletions()[0]
    worker.delete_decision(decision['id'], 'pdf_only')
    assert not client.get('item', 'PARENT23').get('deleted')
    assert not client.get('item', 'NTES2345').get('deleted')
    assert not client.get('item', 'SUPP2345').get('deleted')
    inventory = client.review_action('inventory', {'parent': 'PARENT23'})['inventory']
    assert next(r for r in inventory if r['key'] == 'SURCE234')['deleted']
    worker.sync()
    assert len(worker.state.read('documents')) == 1
    checked('delete_PDF_only_trashes_only_matching_source_and_view_keeps_notes_SI_reference_and_never_resurrects')
    second = cloud.upload_pdf(pdf_bytes(), 'Second validation paper', croot.id, str(uuid.uuid4()), croot.id)
    worker.sync()
    row = next(r for r in worker.review.rows() if r['kind'] == 'cloud' and r['identity'] == second.id)
    worker.approve(row['id'], row['digest'], 'Second.pdf', 'Second validation paper', 'journal_article', 'SYNC2345')
    worker.sync()
    link = worker.state.read('documents')[second.id]
    client.create('item', {'key': 'FULLNT23', 'itemType': 'note', 'parentItem': link['parent'], 'note': '<p>Whole-reference note</p>'})
    client.create('item', {'key': 'FULLSI23', 'itemType': 'attachment', 'parentItem': link['parent'], 'title': 'Whole-reference SI', 'contentType': 'application/pdf', 'linkMode': 'imported_file', 'filename': 'Full SI.pdf'})
    client.upload('FULLSI23', source, '')
    current = next(i for i in cloud.library() if i.id == second.id)
    cloud.update_item(second.id, expected_hash=current.hash, parent='trash', scope_id=croot.id)
    worker.sync()
    snapshot = worker.status()
    decision = next(r for r in worker.review.deletions() if r['payload']['cloud'] == second.id)
    worker.delete_decision(decision['id'], 'everything')
    inventory = client.review_action('inventory', {'parent': link['parent']})['inventory']
    assert next(r for r in inventory if r['key'] == link['parent'])['deleted']
    assert client.get('item', 'FULLNT23')['parentItem'] == link['parent']
    assert client.get('item', 'FULLSI23')['parentItem'] == link['parent']
    checked('delete_everything_uses_native_Trash_for_the_reference_notes_and_SI')
    inbox = root / 'ToSort'; inbox.mkdir()
    worker.config['inbox'] = str(inbox)
    incoming = inbox / 'Download.pdf'; incoming.write_bytes(pdf_bytes())
    from docx import Document
    word = Document(); word.add_heading('Synthetic research notebook', 0); word.add_paragraph('A local intake validation document.'); word.save(inbox / 'Notebook.docx')
    from tests.test_sorting_presentation_extraction import _write_pptx
    _write_pptx(inbox / 'Slides.pptx', ['Local nanoparticle presentation'], title='Local intake slides', author='Example Researcher')
    (inbox / 'Still downloading.pdf').write_bytes(b'%PDF-1.7\nunfinished')
    worker.sync(); time.sleep(2.1); worker.sync()
    rows = [r for r in worker.review.rows() if r['kind'] == 'inbox']
    assert len(rows) == 3
    for row in rows:
        original = Path(row['payload']['path'])
        digest = hashlib.sha256(original.read_bytes()).hexdigest()
        worker.approve(row['id'], row['digest'], 'Reviewed ' + original.name, row['payload']['title'], row['payload']['document_type'], 'SYNC2345')
        finished = worker.review.row(row['id'])
        stored = client.file_path(finished['approval']['source'])
        assert hashlib.sha256(stored.read_bytes()).hexdigest() == digest
        assert not original.exists()
    assert (inbox / 'Still downloading.pdf').exists()
    checked('PDF_Word_and_PowerPoint_intake_waits_for_completion_and_review_then_verifies_managed_copy_before_cleanup')
    from note_organizer.zotero_pilot import reparent_attachment
    source = finished['approval']['source']
    attachment = client.get('item', source)
    old_parent = attachment['parentItem']
    destination = client.create('item', {'key': 'RPAR2345', 'itemType': 'document', 'title': 'Verified reference move'})
    original_bytes = client.file_path(source).read_bytes()
    reparent_attachment(client, source, destination['key'])
    assert client.get('item', source)['parentItem'] == destination['key']
    assert client.file_path(source).read_bytes() == original_bytes
    reparent_attachment(client, source, destination['key'])
    reparent_attachment(client, source, old_parent)
    assert client.get('item', source)['parentItem'] == old_parent
    checked('native_reparent_is_verified_idempotent_and_keeps_attachment_key_and_bytes')
    return snapshot


def bundled_search_checks(client, root, handwriting_pdf, checked):
    """Run local OCR through the plugin's actual private process transport."""
    import pymupdf
    from note_organizer.zotero_pilot import Ledger, Paper, publish
    from note_organizer.zotero_bridge import BridgeState
    pdfs = []
    if handwriting_pdf:
        pdfs.append(('physical-handwriting', handwriting_pdf.resolve(strict=True), 'testing'))
    scanned = root / 'Bundled scan.pdf'
    with pymupdf.open() as printed, pymupdf.open() as image:
        page = printed.new_page(width=400, height=500)
        page.insert_text((35, 100), 'bundledscanzirconium nanoparticle', fontsize=15)
        png = page.get_pixmap(dpi=200).tobytes('png')
        image.new_page(width=400, height=500).insert_image((0, 0, 400, 500), stream=png)
        image.save(scanned)
    pdfs.append(('image-scan', scanned, 'bundledscanzirconium'))
    from note_organizer.zotero_connections import state_directory
    state_home = root / 'profile/notable'
    state = BridgeState(state_directory(state_home, json.loads((state_home / 'config.json').read_text())) / 'bridge')
    links = {}
    for identity, path, keyword in pdfs:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        result = publish(client, Ledger(root / ('search-' + identity)), 'bundled-search-' + identity,
            [Paper(identity, path, digest, {'title': 'Synthetic search validation: ' + identity})], apply=True)['papers'][0]
        key, parent = result['attachment_key'], result['reference_key']
        links[identity] = {'name': 'Synthetic search validation: ' + identity, 'view_key': key, 'parent': parent, 'revision': 'fixture'}
        state.write('documents', links)
        before = client.file_md5(key)
        command(root, 'reindex', key=key)
        assert keyword in client.zot.fulltext_item(key)['content'].lower()
        assert any(r['key'] in {key, parent} for r in client.zot.items(q=keyword, qmode='everything'))
        assert client.file_md5(key) == before and hashlib.sha256(path.read_bytes()).hexdigest() == digest
        checked('packaged_helper_searches_' + identity + '_in_Zotero_without_modifying_PDF')
    # Native desktop comments join the same attachment-content index, while
    # remaining native Zotero annotations rather than being sent to the tablet.
    client.create('item', {'key': 'NATTXT23', 'itemType': 'annotation', 'parentItem': key,
        'annotationType': 'note', 'annotationComment': 'desktopnativequartzkeyword', 'annotationColor': '#ffd400',
        'annotationPageLabel': '1', 'annotationSortIndex': '00000|000000|00000',
        'annotationPosition': json.dumps({'pageIndex': 0, 'rects': [[10, 10, 30, 30]]})})
    command(root, 'reindex', key=key)
    assert any(r['key'] in {key, parent} for r in client.zot.items(q='desktopnativequartzkeyword', qmode='everything'))
    checked('native_Zotero_comments_are_searchable_alongside_print_highlights_and_ink_transcripts')


def bulk_checks(client, root, checked, count, resume=False):
    """Exercise actual native imports, not a mock of Zotero's storage API."""
    from tests.test_zotero_bulk import pdf
    if not resume:
        client.create('collection', {'key': 'BULK2345', 'name': 'Windows imports'})
    source = root / 'Windows literature'; source.mkdir(exist_ok=resume)
    originals = {}
    if resume:
        originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(source.rglob('*.pdf'))
                     if p.name not in {'Broken.pdf', 'Supporting_SI.pdf'}}
        assert len(originals) == count + 1 and len(set(originals.values())) == count
    else:
        for number in range(count):
            path = pdf(source / f'Category {number % 4}' / f'Paper {number:04}.pdf', f'Synthetic bulk scientific study number {number:04}')
            originals[path] = hashlib.sha256(path.read_bytes()).hexdigest()
        duplicate = source / 'Duplicate original.pdf'; duplicate.write_bytes(next(iter(originals)).read_bytes())
        originals[duplicate] = hashlib.sha256(duplicate.read_bytes()).hexdigest()
        (source / 'Broken.pdf').write_bytes(b'not a complete PDF')
        pdf(source / 'Supporting_SI.pdf', 'Supporting information for a synthetic bulk study')
    worker = Worker(root / 'bulk-worker', zotero_factory=lambda: LocalZotero(int(client.endpoint.split(':')[2].split('/')[0]), grant_store=client.grant_store, interactive=False))
    worker.scan_library(str(source), 'BULK2345', lookup=False)
    plan = worker.bulk.read()
    assert len(plan['rows']) == count + 2 and len([r for r in plan['rows'] if r['issues']]) == 2
    if not resume:
        assert not [r for r in client.review_action('catalog')['parents'].values() if 'BULK2345' in r.get('collections', [])]
    checked('native_bulk_scan_is_read_only_groups_exact_duplicates_and_holds_only_broken_and_SI_cases')
    worker.import_library(plan['id'])
    plan = worker.bulk.read()
    assert len([r for r in plan['rows'] if r['status'] == 'complete']) == count, [(r['path'], r['issues']) for r in plan['rows'] if r['issues']]
    def imported_records():
        scope = {'BULK2345'}
        collections = client.collections()
        while True:
            expanded = scope | {r['key'] for r in collections if r.get('parentCollection') in scope}
            if expanded == scope:
                break
            scope = expanded
        return [r for r in client.review_action('catalog')['parents'].values() if set(r.get('collections', [])) & scope]
    records = imported_records()
    categories = [r for r in client.collections() if r.get('parentCollection') == 'BULK2345' and r['name'].startswith('Category ')]
    assert len(categories) == min(4, count)
    assert all(any(r['key'] in parent.get('collections', []) for parent in records) for r in categories)
    checked('native_bulk_import_preserves_original_category_folders')
    if resume:
        catalog = client.review_action('catalog')
        filed = {a['data'].get('parentItem') for a in catalog['attachments'] if a['data'].get('parentItem') in {r['key'] for r in records}}
        assert len(filed) == count and all(r['itemType'] == 'document' for r in records if r['key'] in filed)
        checked('preexisting_reference_only_leftovers_are_retained_without_creating_duplicate_PDFs')
    else:
        assert len(records) == count and all(r['itemType'] == 'document' for r in records)
    for row in plan['rows']:
        if row['status'] == 'complete':
            assert hashlib.sha256(client.file_path(row['attachment']).read_bytes()).hexdigest() == row['digest']
    assert originals == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in originals}
    checked(f'native_bulk_{count}_papers_import_without_complete_citations_or_changes_to_Windows_originals')
    if not resume:
        completed = next(r for r in plan['rows'] if r['status'] == 'complete')
        replayed = client.review_action('bulk_import_status', {'digest': completed['digest'], 'collection': 'BULK2345'})
        assert replayed['attachment'] == completed['attachment']
        replayed_again = client.review_action('bulk_import_status', {'digest': completed['digest'], 'collection': 'BULK2345'})
        assert replayed_again == replayed
        checked('native_import_results_survive_repeated_polling_or_lost_responses')
    worker.scan_library(str(source), 'BULK2345', lookup=False)
    again = worker.bulk.read()
    assert len([r for r in again['rows'] if r['existing']]) == count
    worker.import_library(again['id'])
    assert len(imported_records()) == len(records)
    checked('native_bulk_rescan_reuses_existing_references_and_PDFs_without_creating_duplicates')
    # The normal DOI translator resolves this known preprint without adding a
    # downloaded PDF. The test file contains synthetic text, not a user's paper.
    import pymupdf
    doi_source = root / 'DOI lookup'; doi_source.mkdir()
    known = doi_source / 'Known DOI.pdf'
    paper_title = 'Drug-Eluting, Radiopaque, Tumor-Casting Hydrogels for Endovascular Locoregional Therapy of Hepatocellular Carcinoma'
    with pymupdf.open() as doc:
        page = doc.new_page(); page.insert_text((20, 80), paper_title, fontsize=6)
        page.insert_text((20, 110), 'DOI: 10.1101/2025.11.25.690505', fontsize=10)
        doc.set_metadata({'title': paper_title}); doc.save(known)
    worker.scan_library(str(doi_source), 'BULK2345', lookup=True)
    worker.import_library(worker.bulk.read()['id'])
    row = worker.bulk.read()['rows'][0]
    assert row['status'] == 'complete' and not row['issues'], row
    record = client.get('item', row['parent'])
    assert record['DOI'] == '10.1101/2025.11.25.690505' and record['creators'][0]['lastName'] == 'Dong', record
    assert record['itemType'] == 'preprint' and len(record['creators']) == 9
    assert len(client.children(row['parent'])) == 1
    checked('native_DOI_metadata_keeps_preprint_type_and_author_order_without_downloading_another_PDF')
    # Batch-file routine native imports and match an existing tablet original.
    # Use the same source bytes and in-memory cloud protocol writer as the
    # independently tested bridge; no real tablet account is involved.
    cloud = MemoryCloud()
    croot = cloud.create_folder('Literature', '', str(uuid.uuid4()))
    first = next(iter(originals))
    existing = cloud.upload_pdf(first.read_bytes(), first.stem, croot.id, str(uuid.uuid4()), croot.id)
    worker.cloud_factory = lambda: cloud
    worker.config = {'server_id': client.server_id, 'zotero_root': 'BULK2345', 'tablet_root': croot.id, 'default_folder': croot.id, 'ocr': False}
    worker.sync(); time.sleep(2.1); worker.sync()
    eligible = [r for r in worker.review.rows() if r['kind'] == 'zotero' and r['payload'].get('parent') in {p['key'] for p in records}]
    selected = sorted(eligible, key=lambda r: not any(m['key'] == existing.id for m in r['payload'].get('matches', [])))[:3]
    assert len(selected) == 3
    assert any(any(m['key'] == existing.id for m in r['payload'].get('matches', [])) for r in selected)
    ambiguous = [r for r in selected if not r['payload']['routine']]
    assert len(ambiguous) == 1 and 'several tablet destinations' in str(ambiguous[0]['payload']['issues'])
    held = ambiguous[0]
    worker.approve(held['id'], held['digest'], held['payload']['filename'], held['payload']['title'],
                   held['payload']['document_type'], 'BULK2345', match=existing.id,
                   review_token=held['payload']['review_token'])
    checked('native_import_with_multiple_original_folders_requires_an_explicit_tablet_home')
    worker.approve_routine([{'id': r['id'], 'digest': r['digest'], 'review_token': r['payload']['review_token']}
                            for r in selected if r['payload']['routine']], 'BULK2345')
    assert all(worker.review.row(r['id'])['status'] == 'approved' for r in selected)
    worker.sync()
    assert len(worker.state.read('documents')) == 3 and existing.id in worker.state.read('documents')
    assert len([i for i in cloud.library() if i.is_pdf]) == 3
    checked('native_routine_batch_files_multiple_papers_and_reuses_the_exact_existing_tablet_original')
    tag_id, tag_link = next(iter(worker.state.read('documents').items()))
    key = tag_link['view_key']
    context = worker.tag_context(key, 0)
    ui_tags = command(root, 'reader_tags', context=context)
    command(root, 'snapshot')
    import shutil
    shutil.copyfile(root / 'workspace.png', root / 'tag-workspace.png')
    assert ui_tags['readerButton'] == 'Tag this page' and ui_tags['readerPage'] == 0
    assert ui_tags['requests'][0]['params'] == {'key': key, 'page': 0}
    edit = next(r for r in ui_tags['requests'] if r['method'] == 'set_page_tags')['params']
    worker.dispatch('set_page_tags', edit)
    before_pdf = client.file_path(key).read_bytes()
    worker.sync()
    remote = next(i for i in cloud.library() if i.id == tag_id)
    content = json.loads(cloud.pdf_bundle(remote)[remote.id + '.content'])
    assert any(t['name'] == 'Native paper tag' for t in content['tags'])
    assert any(t['name'] == 'Native page tag' and t['pageId'] == context['page_id'] for t in content['pageTags'])
    assert client.file_path(key).read_bytes() == before_pdf
    assert client.native_annotation_count(key) == 0
    assert worker.tag_search('Native page', 'page')['results'][0]['page_id'] == context['page_id']
    worker.sync()
    assert client.file_path(key).read_bytes() == before_pdf
    checked('actual_reader_tag_button_and_editor_roundtrip_both_scopes_without_creating_annotations_or_rewriting_PDF')
    worker.scan_library(str(source), 'BULK2345', lookup=False)
    assert len([r for r in worker.bulk.read()['rows'] if r['existing']]) == count
    worker.import_library(worker.bulk.read()['id'])
    assert all(r['status'] == 'complete' for r in worker.bulk.read()['rows'] if not r['issues'])
    checked('bulk_rescan_recognizes_originals_after_their_managed_PDFs_receive_tablet_annotations')
    # Real native SDK renders and invokes the private transport too, using a
    # subset already imported above to verify reuse through the packaged UI.
    subset = root / 'UI literature'; subset.mkdir()
    linked = {link['baseline'] for link in worker.state.read('documents').values()}
    unlinked = [p for p in originals if hashlib.sha256(p.read_bytes()).hexdigest() not in linked]
    for number, path in enumerate(unlinked[:3]):
        (subset / f'Original {number}.pdf').write_bytes(path.read_bytes())
    result = command(root, 'rpc', method='scan_library', params={'root': str(subset), 'collection': 'BULK2345', 'lookup': False})['result']
    assert len(result['bulk']['rows']) == 3 and all(r['existing'] for r in result['bulk']['rows'])
    result = command(root, 'rpc', method='import_library', params={'plan_id': result['bulk']['id']})['result']
    assert all(r['status'] == 'complete' for r in result['bulk']['rows'])
    rendered = command(root, 'tab', tab='import')
    assert 'Scan folder — no library changes' in rendered['buttons']
    checked('actual_graphical_import_tab_and_private_worker_scan_and_reuse_existing_PDFs')
    command(root, 'snapshot')
    import shutil
    shutil.copyfile(root / 'workspace.png', root / 'bulk-workspace.png')


def activity_checks(root, checked):
    source = root / 'local-review-preview.pdf'
    source.write_bytes(pdf_bytes())
    snapshot = command(root, 'rpc', method='status')['result']
    snapshot.update(config={'zotero_root': 'ROOT2345', 'tablet_root': 'fixture-tablet', 'automatic': True, 'sync_paused': False},
                    collections=[{'key': 'ROOT2345', 'name': 'Tablet papers', 'parentCollection': False}],
                    mappings={'ROOT2345': {'cloud': 'fixture-tablet'}}, links=[], deletions=[], local_deletions=[],
                    connection={'linked_papers': 0, 'previous': []}, bulk=None,
                    status={'state': 'ready', 'message': 'Ready'})
    snapshot['reviews'] = [{'id': str(number), 'kind': 'cloud', 'identity': str(number), 'status': 'pending', 'digest': 'synthetic',
                            'payload': {'path': str(source), 'title': f'Example paper {number}', 'filename': f'Example {number}.pdf',
                                        'document_type': 'journal_article', 'routine': number == 0, 'details': {}, 'issues': []}}
                           for number in range(2)]
    result = command(root, 'activity_ui', snapshot=snapshot, source=str(source))
    for name in result['passed']:
        checked('native_activity_' + name.replace(' ', '_'))


def tagging_checks(client, root, checked):
    import copy
    import pymupdf
    from note_organizer.zotero_worker import file_sha
    source = root / 'Keyword papers'; source.mkdir()
    path = source / 'Keyword fixture.pdf'
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((40, 80), 'Spectral CT nanoparticle research\n\nIntroduction\nSpectral CT measures iodine concentration. Nanoparticles improve contrast.\nIodine concentration is measured with spectral imaging.')
        document.set_metadata({'keywords': 'Spectral CT\x02nanoparticles'})
        document.save(path)
    digest = file_sha(path)
    collection = client.collections()[0]['key']
    native = client.review_action('bulk_import', {'root': str(source), 'source': str(path), 'digest': digest,
        'md5': hashlib.md5(path.read_bytes()).hexdigest(), 'collection': collection, 'lookup': False,
        'title': 'Keyword fixture paper', 'dois': [], 'evidence': ''})
    while native.get('pending'):
        time.sleep(.2)
        native = client.review_action('bulk_import_status', {'digest': digest, 'collection': collection})
    parent = client.get('item', native['parent']); parent['tags'] = [{'tag': 'Curated label', 'type': 0}]; client.update(parent)
    proposal = command(root, 'rpc', method='suggest_paper_tags', params={'key': native['attachment']})['result']['proposal']
    assert {row['tag'] for row in proposal['keywords']} == {'Spectral CT', 'nanoparticles'}, proposal
    assert [(tag['tag'], tag.get('type', 0)) for tag in client.get('item', native['parent'])['tags']] == [('Curated label', 0)]
    checked('packaged_keyword_analysis_retains_PDF_bytes_and_native_tags_without_writes')
    command(root, 'rpc', method='save_paper_tags', params={'target': proposal['target'], 'token': proposal['token'], 'selected': []})
    tags = client.get('item', native['parent'])['tags']
    assert ('Spectral CT', 1) in [(tag['tag'], tag.get('type', 0)) for tag in tags] and ('Curated label', 0) in [(tag['tag'], tag.get('type', 0)) for tag in tags], tags
    assert not {row['tag'] for row in proposal['suggestions']} & {row['tag'] for row in tags}
    checked('native_keywords_keep_manual_types_and_unselected_suggestions_are_not_added')
    proposal = command(root, 'rpc', method='suggest_paper_tags', params={'key': native['attachment']})['result']['proposal']
    selected = proposal['suggestions'][0]['tag']
    command(root, 'rpc', method='save_paper_tags', params={'target': proposal['target'], 'token': proposal['token'], 'selected': [selected]})
    assert (selected, 0) in [(tag['tag'], tag.get('type', 0)) for tag in client.get('item', native['parent'])['tags']]
    assert file_sha(path) == digest and file_sha(native['path']) == digest
    checked('native_selected_suggestion_is_explicit_and_original_PDF_bytes_are_unchanged')
    snapshot = command(root, 'rpc', method='status')['result']
    snapshot.update(config={'zotero_root': collection, 'tablet_root': 'fixture', 'automatic': False},
        collections=[{'key': collection, 'name': 'Literature', 'parentCollection': False}], mappings={collection: {'cloud': 'fixture'}},
        links=[], deletions=[], local_deletions=[], bulk=None)
    snapshot['reviews'] = [{'id': str(index), 'kind': 'cloud', 'identity': str(index), 'status': 'pending', 'digest': digest,
        'payload': {'path': str(path), 'title': f'Example paper {index:03}', 'filename': f'Example {index:03}.pdf',
                    'document_type': 'journal_article', 'details': {}, 'issues': []}} for index in range(160)]
    proposal = copy.deepcopy(proposal); proposal['selected'] = []
    result = command(root, 'review_list', snapshot=snapshot, proposal=proposal)
    for name in result['passed']:
        checked('native_' + name.replace(' ', '_'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', type=Path, required=True)
    parser.add_argument('--zotero-exe', type=Path, required=True)
    parser.add_argument('--runtime', type=Path)
    parser.add_argument('--xpi', type=Path, help='Test the exact release package instead of rebuilding it')
    parser.add_argument('--handwriting-pdf', type=Path)
    parser.add_argument('--visible', action='store_true')
    parser.add_argument('--keep-open', action='store_true')
    parser.add_argument('--port', type=int, default=23149)
    parser.add_argument('--bulk-count', type=int, default=0)
    parser.add_argument('--skip-ocr', action='store_true', help='Only for development-worker checks without bundled models')
    parser.add_argument('--connection-checks', action='store_true')
    parser.add_argument('--activity-checks', action='store_true')
    parser.add_argument('--tagging-checks', action='store_true')
    parser.add_argument('--existing-fixture', type=Path, help='Copy only a previous synthetic fixture into this new disposable profile')
    args = parser.parse_args()
    root = args.work_dir.resolve(); root.mkdir(parents=True, exist_ok=False)
    profile, data, extension = (root / name for name in ['profile', 'zotero-data', 'test-extension'])
    for path in [profile / 'extensions', data, extension]: path.mkdir(parents=True, exist_ok=True)
    if args.existing_fixture:
        import shutil
        prior = args.existing_fixture.resolve(strict=True)
        record = json.loads((prior / 'test-extension/config.json').read_text())
        if Path(record['data']).resolve() != prior / 'zotero-data' or not (prior / 'done').exists():
            raise ValueError('A closed synthetic fixture is required')
        shutil.copytree(prior / 'zotero-data', data, dirs_exist_ok=True)
        shutil.copytree(prior / 'Windows literature', root / 'Windows literature')
    config = {'data': str(data), 'ready': str(root / 'ready.json'), 'done': str(root / 'done')}
    (extension / 'config.json').write_text(json.dumps(config), encoding='utf-8')
    (extension / 'bootstrap.js').write_text(BOOTSTRAP, encoding='utf-8')
    test_id = 'notable-graphical-test@example.invalid'
    (extension / 'manifest.json').write_text(json.dumps({'manifest_version': 2, 'name': 'Disposable graphical test', 'version': '0.0.1',
        'applications': {'zotero': {'id': test_id, 'strict_min_version': '10.0', 'strict_max_version': '10.*', 'update_url': 'https://example.invalid/updates.json'}}}), encoding='utf-8')
    (profile / 'extensions' / test_id).write_text(str(extension), encoding='utf-8')
    if args.xpi:
        import shutil
        shutil.copyfile(args.xpi.resolve(strict=True), profile / 'extensions/notable-companion@local.invalid.xpi')
    else:
        build(profile / 'extensions/notable-companion@local.invalid.xpi', args.runtime)
    prefs = {'extensions.zotero.firstRun2': False, 'extensions.zotero.firstRunGuidance': False,
        'extensions.zotero.useDataDir': True, 'extensions.zotero.dataDir': str(data), 'extensions.zotero.warnOnUnsafeDataDir': False,
        'extensions.zotero.sync.autoSync': False, 'extensions.zotero.automaticScraperUpdates': False,
        'extensions.zotero.httpServer.enabled': True, 'extensions.zotero.httpServer.port': args.port,
        'extensions.zotero.httpServer.localAPI.enabled': True, 'extensions.zoteroWinWordIntegration.skipInstallation': True,
        'extensions.zoteroOpenOfficeIntegration.skipInstallation': True, 'extensions.update.enabled': False,
        'extensions.enabledScopes': 15, 'extensions.autoDisableScopes': 0, 'extensions.startupScanScopes': 15, 'app.update.enabled': False, 'app.update.auto': False}
    prefs['notable.testDiagnostics'] = True
    if not args.runtime and not args.xpi:
        prefs['notable.developmentWorker'] = json.dumps({'command': sys.executable, 'arguments': [str(Path(__file__).resolve().parent / 'notable_worker.py')]})
    (profile / 'user.js').write_text('\n'.join(f'user_pref({json.dumps(k)}, {json.dumps(v)});' for k, v in prefs.items()), encoding='utf-8')
    startup = subprocess.STARTUPINFO(); startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW; startup.wShowWindow = 1 if args.visible else subprocess.SW_HIDE
    result = {'scope': 'Disposable Zotero profile; no real library or cloud account', 'checks': []}
    def checked(name):
        result['checks'].append({'name': name, 'status': 'PASS'})
        (root / 'results.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    client = None
    with (root / 'stdout.log').open('w') as stdout, (root / 'stderr.log').open('w') as stderr:
        process = subprocess.Popen([str(args.zotero_exe.resolve(strict=True)), '-no-remote', '-profile', str(profile), '-datadir', str(data), *([] if args.visible else ['-headless'])], startupinfo=startup, stdout=stdout, stderr=stderr)
        try:
            deadline = time.monotonic() + 180
            while not Path(config['ready']).exists():
                # A normal first launch may restart after discovering an XPI.
                # The scoped ready file, rather than the launcher PID, is the
                # authority for the instance that actually owns this profile.
                if time.monotonic() > deadline: raise RuntimeError('Disposable Zotero failed to start')
                time.sleep(.2)
            from note_organizer.remarkable_cloud import TokenVault
            client = LocalZotero(args.port, grant_store=TokenVault(profile / 'notable/zotero-grant.dpapi'))
            assert client.server_id == json.loads(Path(config['ready']).read_text())['server']
            ui = command(root, 'open')
            assert ui['worker'] and any(name in ui['title'] for name in ['Notable', 'reMarkable Sync']) and 'Connect Zotero' in ui['buttons'], ui
            checked('Zotero_opens_graphical_setup_and_manages_private_worker')
            command(root, 'connect')
            command(root, 'tab', tab='review')
            checked('graphical_connect_persists_permission_without_terminal')
            snapshot = native_checks(client, root, checked) if not args.existing_fixture else None
            if args.bulk_count:
                bulk_checks(client, root, checked, args.bulk_count, resume=bool(args.existing_fixture))
            if args.connection_checks:
                connection_checks(client, root, checked)
            if args.activity_checks:
                activity_checks(root, checked)
            if args.tagging_checks:
                tagging_checks(client, root, checked)
            if not args.skip_ocr:
                bundled_search_checks(client, root, args.handwriting_pdf, checked)
            if snapshot:
                rendered = command(root, 'render', snapshot=snapshot)
                assert all(label in rendered['buttons'] for label in ['Restore on tablet', 'Delete PDF only', 'Delete everything'])
                checked('actual_review_window_renders_all_three_deletion_choices_and_new_documents')
            command(root, 'snapshot')
            (root / 'results.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            if args.keep_open:
                print(json.dumps({'state': 'ready_for_visual_inspection', 'root': str(root)}), flush=True)
                while not (root / 'close-test').exists(): time.sleep(1)
        except Exception as error:
            result['error'] = str(error)
            raise
        finally:
            if client: client.close()
            Path(config['done']).touch()
            try: process.wait(timeout=15)
            except subprocess.TimeoutExpired: process.terminate(); process.wait(timeout=10)
            (root / 'results.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
            print(json.dumps(result, indent=2))


if __name__ == '__main__': main()
